#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generador de fichas de inscripción — CEPA Suroccidente
=======================================================

Rellena automáticamente el impreso de inscripción (Word) para cada
alumno/a de un listado en Excel (el generado por el formulario de
pre-matrícula) y guarda un documento .docx por persona.

No todos los campos del impreso existen en el Excel: los que faltan
se dejan en blanco para rellenar a mano.

Requisitos (una sola vez, desde una consola):
    pip install python-docx openpyxl

Uso:
    Doble clic sobre este archivo (generador_inscripciones.pyw)
    - Se abre sin ventana de consola (extensión .pyw).
"""

import os
import re
import sys
import queue
import unicodedata
import threading
import subprocess
import traceback
from datetime import datetime

import tkinter as tk
from tkinter import ttk, filedialog, messagebox

# ---------------------------------------------------------------------------
# Comprobación de dependencias
# ---------------------------------------------------------------------------
MISSING = []
try:
    import openpyxl
except ImportError:
    MISSING.append("openpyxl")

try:
    import docx
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    MISSING.append("python-docx")

if MISSING:
    import tkinter.messagebox as mb
    root = tk.Tk()
    root.withdraw()
    mb.showerror(
        "Faltan librerías",
        "Este programa necesita las siguientes librerías de Python, que no "
        "están instaladas:\n\n"
        + "\n".join(f"  • {m}" for m in MISSING)
        + "\n\nAbre una consola (CMD) y ejecuta:\n\n"
        f"    pip install {' '.join(MISSING)}\n\n"
        "y vuelve a abrir este programa."
    )
    sys.exit(1)


# ---------------------------------------------------------------------------
# Utilidades de texto
# ---------------------------------------------------------------------------
def normaliza(texto):
    """minúsculas, sin acentos, sin espacios sobrantes."""
    if texto is None:
        return ""
    texto = str(texto).strip().lower()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"\s+", " ", texto)
    return texto


def limpia_nombre_archivo(texto):
    """Quita caracteres no válidos para nombres de archivo en Windows."""
    texto = (texto or "").strip()
    texto = re.sub(r'[\\/:*?"<>|]', "", texto)
    texto = re.sub(r"\s+", " ", texto).strip()
    return texto or "sin_nombre"


def valor_texto(v):
    """Convierte el valor de una celda de Excel a texto legible."""
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


# ---------------------------------------------------------------------------
# Localización flexible de columnas del Excel (por nombre de cabecera)
# ---------------------------------------------------------------------------
# Cada campo interno se busca por una lista de palabras clave que deben
# aparecer todas (sin acentos) en el encabezado de la columna.
CAMPOS_EXCEL = {
    "id":                  [["id"]],
    "nombre":              [["nombre"]],
    "apellidos":           [["apellido"]],
    "dni":                 [["dni"], ["nie"], ["pasaporte"]],
    "telefono":            [["telefono"], ["movil"]],
    "aula":                [["aula"]],
    "ensenanza":           [["ensenanza"]],
    "nacionalidad":        [["nacionalidad"]],
    "fecha_nacimiento":    [["fecha", "nacimiento"]],
    "localidad_nacimiento":[["localidad", "nacimiento"]],
    "pais_nacimiento":     [["pais", "nacimiento"]],
    "residencia":          [["residencia"]],
    "direccion":           [["direccion"]],
    "cp":                  [["codigo", "postal"], ["cp"]],
    "provincia":           [["provincia"]],
    "correo":              [["correo"], ["email"], ["e-mail"]],
    "estudios":            [["estudios"]],
    "hora_inicio":         [["hora", "inicio"]],
}

MESES_ES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
            "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def localizar_columnas(cabeceras):
    """
    cabeceras: lista de textos de la fila 1 del Excel (en orden de columna).
    Devuelve dict campo_interno -> índice de columna (0-based), o None si
    no se encuentra esa columna en el Excel.
    """
    normalizadas = [normaliza(h) for h in cabeceras]
    resultado = {}
    for campo, listas_claves in CAMPOS_EXCEL.items():
        encontrado = None
        for idx, texto in enumerate(normalizadas):
            for claves in listas_claves:
                if all(clave in texto for clave in claves):
                    encontrado = idx
                    break
            if encontrado is not None:
                break
        resultado[campo] = encontrado
    return resultado


# ---------------------------------------------------------------------------
# Interpretación de la columna "Enseñanza y turno"
# ---------------------------------------------------------------------------
# Programas que aparecen como valor único, sin turno asociado.
PROGRAMAS_SIMPLES = {
    "espad 1.2": {"row": 5, "check_col": 1},
    "espad 2.1": {"row": 5, "check_col": 5},
    "espad 2.2": {"row": 5, "check_col": 7},
    "ptgeso": {"row": 7, "check_col": 8, "aula_col": 9},
    "grado superior": {"row": 9, "check_col": 8, "aula_col": 9},
    "ciclos grado superior": {"row": 9, "check_col": 8, "aula_col": 9},
    "formacion basica i": {"row": 1, "check_col": 8, "aula_col": 9},
    "formacion basica ii": {"row": 2, "check_col": 8, "aula_col": 9},
    "competencias basicas 3": {"row": 3, "check_col": 8, "aula_col": 9},
    "competencias basicas 4": {"row": 4, "check_col": 8, "aula_col": 9},
}

# Materias que se combinan con un turno (MAÑANA / TARDE).
MATERIAS_CON_TURNO = {
    "espanol": {"row": 12, "check_col": 8, "aula_col": 9, "manana_col": 11, "tarde_col": 13},
    "ingles": {"row": 13, "check_col": 8, "aula_col": 9, "manana_col": 11, "tarde_col": 13},
    "informatica": {"row": 14, "check_col": 8, "aula_col": 9, "manana_col": 11, "tarde_col": 13},
}

RE_TURNO = re.compile(r"^(.*?)\s+(MA[ÑN]ANAS?|TARDES?)$", re.IGNORECASE)


def interpreta_ensenanzas(texto, avisos):
    """
    Convierte el texto de la columna "Enseñanza y turno" (p.ej.
    "Inglés TARDES;Informática MAÑANAS;") en una lista de selecciones:
        [{"row":.., "check_col":.., "aula_col":.., "turno_col":..}, ...]
    """
    selecciones = []
    if not texto:
        return selecciones
    tokens = [t.strip() for t in str(texto).split(";") if t.strip()]
    for token in tokens:
        clave = normaliza(token)
        if clave in PROGRAMAS_SIMPLES:
            info = dict(PROGRAMAS_SIMPLES[clave])
            info["turno_col"] = None
            selecciones.append(info)
            continue

        m = RE_TURNO.match(token)
        if m:
            materia_txt, turno_txt = m.group(1), m.group(2)
            materia_norm = normaliza(materia_txt)
            turno_norm = normaliza(turno_txt)
            if materia_norm in MATERIAS_CON_TURNO:
                base = MATERIAS_CON_TURNO[materia_norm]
                info = dict(row=base["row"], check_col=base["check_col"],
                            aula_col=base["aula_col"])
                if turno_norm.startswith("manana"):
                    info["turno_col"] = base["manana_col"]
                else:
                    info["turno_col"] = base["tarde_col"]
                selecciones.append(info)
                continue

        avisos.append(f'No se reconoce la enseñanza "{token}" (se deja sin marcar).')
    return selecciones


# ---------------------------------------------------------------------------
# Relleno del documento Word
# ---------------------------------------------------------------------------
def escribe_celda(cell, texto, negrita=False, tamano=10, centrado=False):
    if texto is None:
        return
    texto = str(texto).strip()
    # Vacía la celda conservando el primer párrafo
    for p in cell.paragraphs[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for run in list(p.runs):
        run._element.getparent().remove(run._element)
    if not texto:
        return
    run = p.add_run(texto)
    run.font.size = Pt(tamano)
    run.font.bold = negrita
    run.font.name = "Calibri"
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def reemplaza_curso(document, curso_texto):
    """Sustituye 'CURSO 20__ /20__' por el curso académico indicado."""
    for p in document.paragraphs[:8]:
        if "CURSO" in p.text.upper() and "20" in p.text:
            texto_completo = p.text
            nuevo = re.sub(r"20_+\s*/\s*20_+", curso_texto, texto_completo)
            if nuevo != texto_completo:
                for run in list(p.runs):
                    run._element.getparent().remove(run._element)
                r = p.add_run(nuevo)
                r.font.bold = True
            return


PATRON_FECHA_FIRMA = re.compile(r"a\s+\.+\s*de\s+\.+\s*de\s+20_+\s*_*", re.IGNORECASE)


def escribe_fecha_firma(document, fecha_texto, avisos):
    """
    Rellena la línea "En ..., a ... de ... de 20__" con la fecha de
    la columna "Hora de inicio" del Excel (día y mes en letra + año).
    fecha_texto se espera en formato dd/mm/aaaa.
    """
    if not fecha_texto:
        return
    try:
        fecha = datetime.strptime(fecha_texto.strip(), "%d/%m/%Y")
    except ValueError:
        avisos.append(f'No se pudo interpretar la fecha "{fecha_texto}" para la firma.')
        return

    dia = fecha.day
    mes = MESES_ES[fecha.month - 1]
    anio = fecha.year
    reemplazo = f"a {dia} de {mes} de {anio}"

    for p in document.paragraphs:
        texto = p.text
        if "20_" not in texto:
            continue
        nuevo = PATRON_FECHA_FIRMA.sub(reemplazo, texto)
        if nuevo != texto:
            for run in list(p.runs):
                run._element.getparent().remove(run._element)
            p.add_run(nuevo)
            return

    avisos.append("No se encontró la línea de fecha de firma en la plantilla.")


def rellena_documento(plantilla_path, datos, curso_texto, avisos):
    """
    datos: dict con las claves de CAMPOS_EXCEL (algunas pueden faltar/None)
    Devuelve un objeto Document ya relleno.
    """
    document = docx.Document(plantilla_path)

    if curso_texto:
        reemplaza_curso(document, curso_texto)

    escribe_fecha_firma(document, datos.get("hora_inicio", ""), avisos)

    tablas = document.tables
    # --- Tabla 0: Nº de solicitud -------------------------------------
    try:
        t0 = tablas[0]
        if datos.get("id"):
            escribe_celda(t0.rows[1].cells[3], datos["id"], centrado=True)
    except Exception:
        avisos.append("No se pudo escribir el número de solicitud.")

    # --- Tabla 1: Datos personales -------------------------------------
    try:
        t1 = tablas[1]
        escribe_celda(t1.rows[1].cells[0], datos.get("nombre", ""))
        escribe_celda(t1.rows[1].cells[1], datos.get("apellidos", ""))
        escribe_celda(t1.rows[1].cells[6], datos.get("dni", ""))

        escribe_celda(t1.rows[3].cells[0], datos.get("telefono", ""))
        escribe_celda(t1.rows[3].cells[1], datos.get("fecha_nacimiento", ""))
        escribe_celda(t1.rows[3].cells[2], datos.get("localidad_nacimiento", ""))
        escribe_celda(t1.rows[3].cells[5], datos.get("pais_nacimiento", ""))

        escribe_celda(t1.rows[5].cells[0], datos.get("direccion", ""))
        escribe_celda(t1.rows[5].cells[3], datos.get("cp", ""))
        escribe_celda(t1.rows[5].cells[4], datos.get("residencia", ""))
        escribe_celda(t1.rows[5].cells[6], datos.get("provincia", ""))

        escribe_celda(t1.rows[7].cells[0], datos.get("correo", ""))
        escribe_celda(t1.rows[7].cells[3], datos.get("nacionalidad", ""))
        escribe_celda(t1.rows[7].cells[4], datos.get("estudios", ""))
    except Exception:
        avisos.append("No se pudieron escribir todos los datos personales.")

    # --- Tabla 2: Enseñanzas -------------------------------------------
    try:
        t2 = tablas[2]
        selecciones = interpreta_ensenanzas(datos.get("ensenanza", ""), avisos)
        for sel in selecciones:
            fila = t2.rows[sel["row"]]
            escribe_celda(fila.cells[sel["check_col"]], "X", negrita=True, centrado=True)
            if sel.get("aula_col") is not None and datos.get("aula"):
                escribe_celda(fila.cells[sel["aula_col"]], datos["aula"], centrado=True)
            if sel.get("turno_col") is not None:
                escribe_celda(fila.cells[sel["turno_col"]], "X", negrita=True, centrado=True)
    except Exception:
        avisos.append("No se pudieron marcar las enseñanzas solicitadas.")

    return document


# ---------------------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------------------
def lee_alumnos(excel_path, avisos):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    filas = list(ws.iter_rows(values_only=True))
    if not filas:
        raise ValueError("El Excel está vacío.")

    cabeceras = list(filas[0])
    columnas = localizar_columnas(cabeceras)

    if columnas.get("nombre") is None or columnas.get("apellidos") is None:
        raise ValueError(
            "No se encuentran las columnas de Nombre y/o Apellidos en el Excel."
        )

    faltantes = [c for c, idx in columnas.items() if idx is None]
    if faltantes:
        avisos.append(
            "Columnas no encontradas en el Excel (se dejarán en blanco): "
            + ", ".join(faltantes)
        )

    alumnos = []
    for fila in filas[1:]:
        if fila is None or all(v is None for v in fila):
            continue

        def obtiene(campo):
            idx = columnas.get(campo)
            if idx is None or idx >= len(fila):
                return ""
            return valor_texto(fila[idx])

        nombre = obtiene("nombre")
        apellidos = obtiene("apellidos")
        if not nombre and not apellidos:
            continue

        datos = {campo: obtiene(campo) for campo in CAMPOS_EXCEL}
        alumnos.append(datos)

    return alumnos


# ---------------------------------------------------------------------------
# Proceso completo
# ---------------------------------------------------------------------------
def genera_documentos(excel_path, plantilla_path, carpeta_destino, curso_texto, log_fn, progreso_fn):
    avisos_generales = []
    alumnos = lee_alumnos(excel_path, avisos_generales)
    for aviso in avisos_generales:
        log_fn(f"⚠ {aviso}")

    if not alumnos:
        raise ValueError("No se ha encontrado ningún alumno en el Excel.")

    os.makedirs(carpeta_destino, exist_ok=True)

    total = len(alumnos)
    generados = 0
    nombres_usados = {}

    for i, datos in enumerate(alumnos, start=1):
        avisos = []
        nombre = datos.get("nombre", "").strip()
        apellidos = datos.get("apellidos", "").strip()
        etiqueta = f"{apellidos} {nombre}".strip() or f"alumno_{i}"

        try:
            document = rellena_documento(plantilla_path, datos, curso_texto, avisos)

            base = limpia_nombre_archivo(f"{apellidos} {nombre}".strip())
            nombre_archivo = f"{base}.docx"
            contador = nombres_usados.get(base, 0)
            if contador:
                nombre_archivo = f"{base} ({contador + 1}).docx"
            nombres_usados[base] = contador + 1

            destino = os.path.join(carpeta_destino, nombre_archivo)
            document.save(destino)
            generados += 1
            log_fn(f"✓ {etiqueta}  →  {nombre_archivo}")
            for a in avisos:
                log_fn(f"    ⚠ {a}")
        except Exception as e:
            log_fn(f"✗ {etiqueta}: ERROR — {e}")
            log_fn("    " + traceback.format_exc(limit=1).strip().splitlines()[-1])

        progreso_fn(i, total)

    return generados, total


# ---------------------------------------------------------------------------
# Interfaz gráfica
# ---------------------------------------------------------------------------
class App(tk.Tk):
    PAD = 8

    def __init__(self):
        super().__init__()
        self.title("Generador de fichas de inscripción — CEPA Suroccidente")
        self.resizable(False, False)
        self.configure(bg="#f4f6f8")

        self.excel_path = tk.StringVar()
        self.plantilla_path = tk.StringVar()
        self.destino_path = tk.StringVar()
        self.curso_var = tk.StringVar(value=self._curso_por_defecto())
        self.abrir_carpeta_var = tk.BooleanVar(value=True)

        self._cola = queue.Queue()
        self._hilo_en_marcha = False

        self._construye_interfaz()
        self._centra_ventana()
        self.after(100, self._procesa_cola)

    # ------------------------------------------------------------------
    def _curso_por_defecto(self):
        hoy = datetime.now()
        ini = hoy.year if hoy.month >= 8 else hoy.year - 1
        return f"{ini}/{ini + 1}"

    def _centra_ventana(self):
        self.update_idletasks()
        ancho = self.winfo_reqwidth()
        alto = self.winfo_reqheight()
        pantalla_ancho = self.winfo_screenwidth()
        x = max((pantalla_ancho - ancho) // 2, 0)
        y = 5  # 5 píxeles del borde superior
        self.geometry(f"{ancho}x{alto}+{x}+{y}")

    # ------------------------------------------------------------------
    def _construye_interfaz(self):
        estilo = ttk.Style(self)
        try:
            estilo.theme_use("clam")
        except tk.TclError:
            pass
        estilo.configure("TLabel", background="#f4f6f8", font=("Segoe UI", 9))
        estilo.configure("Titulo.TLabel", background="#f4f6f8", font=("Segoe UI", 13, "bold"))
        estilo.configure("Sub.TLabel", background="#f4f6f8", foreground="#555555", font=("Segoe UI", 8))
        estilo.configure("TButton", font=("Segoe UI", 9))
        estilo.configure("Generar.TButton", font=("Segoe UI", 10, "bold"))
        estilo.configure("TCheckbutton", background="#f4f6f8", font=("Segoe UI", 9))

        contenedor = ttk.Frame(self, padding=16, style="TFrame")
        estilo.configure("TFrame", background="#f4f6f8")
        contenedor.grid(row=0, column=0, sticky="nsew")

        ttk.Label(contenedor, text="Generador de fichas de inscripción",
                  style="Titulo.TLabel").grid(row=0, column=0, columnspan=3, sticky="w")
        ttk.Label(contenedor, text="CEPA Suroccidente — a partir del listado de pre-matrícula",
                  style="Sub.TLabel").grid(row=1, column=0, columnspan=3, sticky="w", pady=(0, 14))

        self._fila_archivo(contenedor, 2, "Listado de alumnos (Excel):",
                            self.excel_path, self._elige_excel)
        self._fila_archivo(contenedor, 3, "Plantilla del impreso (Word):",
                            self.plantilla_path, self._elige_plantilla)
        self._fila_archivo(contenedor, 4, "Carpeta de destino:",
                            self.destino_path, self._elige_destino, es_carpeta=True)

        ttk.Label(contenedor, text="Curso académico:").grid(row=5, column=0, sticky="w", pady=(10, 0))
        ttk.Entry(contenedor, textvariable=self.curso_var, width=14).grid(
            row=5, column=1, sticky="w", pady=(10, 0))

        ttk.Checkbutton(contenedor, text="Abrir la carpeta de destino al terminar",
                         variable=self.abrir_carpeta_var).grid(
            row=6, column=0, columnspan=3, sticky="w", pady=(8, 4))

        self.boton_generar = ttk.Button(contenedor, text="Generar documentos",
                                         style="Generar.TButton", command=self._on_generar)
        self.boton_generar.grid(row=7, column=0, columnspan=3, sticky="ew", pady=(10, 10))

        self.barra = ttk.Progressbar(contenedor, orient="horizontal", mode="determinate")
        self.barra.grid(row=8, column=0, columnspan=3, sticky="ew")

        self.etiqueta_estado = ttk.Label(contenedor, text="Listo.")
        self.etiqueta_estado.grid(row=9, column=0, columnspan=3, sticky="w", pady=(4, 6))

        marco_log = ttk.Frame(contenedor)
        marco_log.grid(row=10, column=0, columnspan=3, sticky="nsew")
        self.texto_log = tk.Text(marco_log, width=76, height=14, font=("Consolas", 9),
                                  state="disabled", bg="white", relief="solid", borderwidth=1)
        scroll = ttk.Scrollbar(marco_log, orient="vertical", command=self.texto_log.yview)
        self.texto_log.configure(yscrollcommand=scroll.set)
        self.texto_log.grid(row=0, column=0, sticky="nsew")
        scroll.grid(row=0, column=1, sticky="ns")

        contenedor.columnconfigure(1, weight=1)

    def _fila_archivo(self, padre, fila, etiqueta, variable, comando, es_carpeta=False):
        ttk.Label(padre, text=etiqueta).grid(row=fila, column=0, sticky="w", pady=3)
        entrada = ttk.Entry(padre, textvariable=variable, width=52)
        entrada.grid(row=fila, column=1, sticky="ew", padx=(6, 6), pady=3)
        ttk.Button(padre, text="Examinar…", command=comando).grid(row=fila, column=2, pady=3)

    # ------------------------------------------------------------------
    def _elige_excel(self):
        ruta = filedialog.askopenfilename(
            title="Selecciona el Excel del listado de alumnos",
            filetypes=[("Excel", "*.xlsx *.xlsm"), ("Todos los archivos", "*.*")])
        if ruta:
            self.excel_path.set(ruta)

    def _elige_plantilla(self):
        ruta = filedialog.askopenfilename(
            title="Selecciona la plantilla del impreso (Word)",
            filetypes=[("Word", "*.docx"), ("Todos los archivos", "*.*")])
        if ruta:
            self.plantilla_path.set(ruta)

    def _elige_destino(self):
        ruta = filedialog.askdirectory(title="Selecciona la carpeta donde guardar los documentos")
        if ruta:
            self.destino_path.set(ruta)

    # ------------------------------------------------------------------
    def _log(self, texto):
        self._cola.put(("log", texto))

    def _progreso(self, actual, total):
        self._cola.put(("progreso", (actual, total)))

    def _procesa_cola(self):
        try:
            while True:
                tipo, valor = self._cola.get_nowait()
                if tipo == "log":
                    self.texto_log.configure(state="normal")
                    self.texto_log.insert("end", valor + "\n")
                    self.texto_log.see("end")
                    self.texto_log.configure(state="disabled")
                elif tipo == "progreso":
                    actual, total = valor
                    self.barra["maximum"] = total
                    self.barra["value"] = actual
                    self.etiqueta_estado.configure(text=f"Procesando {actual} de {total}…")
                elif tipo == "fin":
                    self._hilo_en_marcha = False
                    self.boton_generar.configure(state="normal")
                    generados, total, carpeta, error = valor
                    if error:
                        self.etiqueta_estado.configure(text="Ha ocurrido un error.")
                        messagebox.showerror("Error", str(error))
                    else:
                        self.etiqueta_estado.configure(
                            text=f"Terminado: {generados} de {total} documentos generados.")
                        messagebox.showinfo(
                            "Proceso terminado",
                            f"Se han generado {generados} de {total} documentos en:\n{carpeta}")
                        if self.abrir_carpeta_var.get():
                            self._abre_carpeta(carpeta)
        except queue.Empty:
            pass
        self.after(100, self._procesa_cola)

    # ------------------------------------------------------------------
    def _abre_carpeta(self, carpeta):
        try:
            if sys.platform.startswith("win"):
                os.startfile(carpeta)  # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                subprocess.Popen(["open", carpeta])
            else:
                subprocess.Popen(["xdg-open", carpeta])
        except Exception as e:
            self._log(f"⚠ No se pudo abrir la carpeta automáticamente: {e}")

    # ------------------------------------------------------------------
    def _on_generar(self):
        if self._hilo_en_marcha:
            return

        excel_path = self.excel_path.get().strip()
        plantilla_path = self.plantilla_path.get().strip()
        destino = self.destino_path.get().strip()
        curso = self.curso_var.get().strip()

        if not excel_path or not os.path.isfile(excel_path):
            messagebox.showwarning("Falta información", "Selecciona el archivo Excel del listado.")
            return
        if not plantilla_path or not os.path.isfile(plantilla_path):
            messagebox.showwarning("Falta información", "Selecciona la plantilla Word del impreso.")
            return
        if not destino:
            messagebox.showwarning("Falta información", "Selecciona la carpeta de destino.")
            return

        self.texto_log.configure(state="normal")
        self.texto_log.delete("1.0", "end")
        self.texto_log.configure(state="disabled")
        self.barra["value"] = 0
        self.etiqueta_estado.configure(text="Procesando…")
        self.boton_generar.configure(state="disabled")
        self._hilo_en_marcha = True

        hilo = threading.Thread(
            target=self._trabajo_en_segundo_plano,
            args=(excel_path, plantilla_path, destino, curso),
            daemon=True,
        )
        hilo.start()

    def _trabajo_en_segundo_plano(self, excel_path, plantilla_path, destino, curso):
        try:
            generados, total = genera_documentos(
                excel_path, plantilla_path, destino, curso,
                log_fn=self._log, progreso_fn=self._progreso)
            self._cola.put(("fin", (generados, total, destino, None)))
        except Exception as e:
            self._log("✗ ERROR GENERAL: " + str(e))
            self._log(traceback.format_exc())
            self._cola.put(("fin", (0, 0, destino, e)))


if __name__ == "__main__":
    app = App()
    app.mainloop()
