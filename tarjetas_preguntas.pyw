# -*- coding: utf-8 -*-
import random
import os
import subprocess
import platform
import tkinter as tk
from tkinter import filedialog, messagebox
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Banco inicial de 100 preguntas por defecto
BANCO_PREGUNTAS_DEFAULT = [
    # Infancia y primeros años
    "¿Cuál es el primer recuerdo que guardas de cuando eras niño?",
    "¿Cómo era tu casa familiar y qué rincón era tu preferido?",
    "¿A qué jugabas en la calle con tus amigos del barrio?",
    "¿Recuerdas alguna travesura de la infancia por la que te regañaran?",
    "¿Cuál era tu plato de comida favorito cuando eras pequeño?",
    "¿Había algún alimento que te negaras a comer en esa época?",
    "¿Cómo celebrabais los cumpleaños en tu familia?",
    "¿Qué ropa o uniforme recuerdas haber llevado con más desgana?",
    "¿Tienes alguna cicatriz de la infancia que guarde una historia?",
    "¿Cuál era el apodo o mote cariñoso que te ponían de niño?",
    "¿Qué persona fuera de tu familia inmediata marcó tu infancia?",
    "¿Cuál fue el regalo de Reyes o cumpleaños que más ilusión te hizo de niño?",
    "¿Qué canciones o coplas tarareabas o escuchabas en casa de pequeño?",
    "¿Cómo viviste la Navidad o las fiestas de tu pueblo en tu niñez?",
    "¿Ayudabas en las tareas de la casa o en algún pequeño trabajo de niño?",
    
    # Juventud y adolescencia
    "¿Cómo era tu grupo de amigos de la adolescencia y qué solíais hacer?",
    "¿Qué música o artistas sonaban en la radio cuando eras joven?",
    "¿Cuál fue tu primer baile o la primera vez que fuiste a una verbena?",
    "¿Qué ropa o estilo de moda era el que más te gustaba llevar de joven?",
    "¿Recuerdas alguna escapada o excursión de juventud con amigos?",
    "¿Qué querías ser de mayor cuando terminaste el colegio?",
    "¿Cómo fue tu primer día en el instituto o la escuela superior?",
    "¿Cuál era tu asignatura favorita y cuál se te hacía más cuesta arriba?",
    "¿Te costó mucho aprobar el carné de conducir o aprender a llevar algún vehículo?",
    "¿Cuál fue tu primer sueldo o el dinero que ganaste por tu cuenta por primera vez?",
    "¿Qué afición o deporte practicabas con más entusiasmo en esa época?",
    "¿Fuiste a ver algún concierto o espectáculo que te impresionara mucho de joven?",
    "¿Cuál fue el primer libro, tebeo o revista que leíste por puro placer?",
    "¿Qué lugar o bar del barrio o pueblo era vuestro punto de reunión habitual?",
    "¿Recuerdas cómo viviste tus primeros años de independencia o madurez?",

    # El amor, la pareja y la familia
    "¿Cómo y dónde conociste a tu mujer o a tu gran amor?",
    "¿Qué fue lo primero que te llamó la atención de ella cuando la viste?",
    "¿Cómo fue la pedida de mano o el momento en que decidisteis formar una vida juntos?",
    "¿Qué anécdota recuerdas del día de vuestra boda?",
    "¿Cuál ha sido el secreto, según tú, para mantener una relación tantos años?",
    "¿Cómo recuerdas el día en que nació vuestro primer hijo?",
    "¿Qué nombre pensasteis para vuestros hijos y por qué los elegisteis?",
    "¿Qué travesuras de tus hijos te sacaban de quicio pero hoy recuerdas con risa?",
    "¿Cómo era un domingo típico en familia cuando los hijos eran pequeños?",
    "¿Cuál ha sido el viaje familiar que más disfrutasteis todos juntos?",
    "¿Qué consejo de tus padres has repetido o aplicado con tus propios hijos?",
    "¿Cómo viviste el momento en que los hijos empezaron a volar del nido?",
    "¿Qué se siente al ser abuelo por primera vez?",
    "¿Qué te gusta hacer ahora con tus nietos o qué te gustaba hacer cuando eran pequeños?",
    "¿Cuál es el momento más feliz que guardas en la memoria junto a tu familia?",

    # Vida profesional y anécdotas laborales
    "¿Cómo conseguiste tu primer trabajo formal?",
    "¿Cómo era tu jefe o tus compañeros en aquel primer empleo?",
    "¿Cuál era tu rutina diaria para ir a trabajar?",
    "¿Qué es lo que más te gustaba de tu trabajo y qué lo que menos?",
    "¿Recuerdas alguna anécdota divertida o un gran apuro que pasaras en el trabajo?",
    "¿Hiciste buenos amigos en el ámbito laboral que conservaras con los años?",
    "¿Cómo cambió tu trabajo desde que empezaste hasta que te jubilaste?",
    "¿Tuviste que hacer algún sacrificio importante por motivos laborales?",
    "¿Qué sentiste el último día antes de jubilarte?",
    "¿Crees que hoy en día se trabaja de forma muy distinta a tu época?",

    # El mundo, los viajes y los cambios de época
    "¿Cuál fue el primer viaje largo que hiciste fuera de tu provincia o región?",
    "¿Qué tecnología o invento te ha parecido el más asombroso de toda tu vida?",
    "¿Cómo cambió tu vida cuando llegó la televisión a casa por primera vez?",
    "¿Recuerdas cómo se vivía antes sin teléfonos móviles ni internet?",
    "¿Qué lugar del mundo te habría gustado visitar y no pudiste?",
    "¿Cuál es el sitio al que siempre te gusta volver de vacaciones?",
    "¿Cómo viviste los cambios políticos o sociales importantes de tu juventud y madurez?",
    "¿Qué costumbres de antes te da pena que se hayan perdido hoy en día?",
    "¿Qué medio de transporte (tren, barco, coche, avión) recuerdas con más emoción?",
    "¿Cómo eran los veranos de antes en comparación con los de ahora?",

    # Gustos, aficiones y pasiones
    "¿Cuál ha sido tu pasatiempo o hobby favorito a lo largo de los años?",
    "¿Qué tipo de música o género de canciones prefieres escuchar?",
    "¿Eres de leer libros, periódicos, o prefieres ver la televisión y documentales?",
    "¿Practicaste algún deporte o fuiste muy aficionado a ver alguno en la televisión?",
    "¿Se te da bien cocinar algún plato estrella? ¿Cuál es tu especialidad?",
    "¿Tienes o tuviste alguna mascota a la que recuerdes con especial cariño?",
    "¿Qué prefieres: el campo y la montaña o la playa y el mar?",
    "¿Cuál es tu bebida o aperitivo favorito para un día de celebración?",
    "¿Qué juego de mesa o de cartas (dominó, mus, brisca) se te da mejor?",
    "¿Eres de los que prefiere madrugar o aprovechar la noche?",

    # Personas inolvidables y anécdotas
    "¿Qué profesor de la escuela o el colegio te marcó más y por qué no lo has olvidado?",
    "¿De qué amigo de la infancia o juventud te gustaría saber algo hoy en día?",
    "¿Quién ha sido el personaje famoso, artista o político que más te ha impresionado conocer o ver?",
    "¿Qué vecino de los de toda la vida guardas con mejor recuerdo?",
    "¿Te encontraste alguna vez con alguien famoso por la calle de casualidad?",
    "¿Quién fue la persona más sabia o con más recursos que conociste en tu vida?",
    "¿Qué vecino o compañero de mili (si la hizo) te dio los momentos más divertidos?",
    "¿Qué persona te tendió una mano cuando más lo necesitabas?",
    "¿Hay alguien a quien te hubiera gustado pedirle perdón o darle las gracias?",
    "¿Quién te enseñó el oficio o la lección de vida más importante que tienes?",

    # Reflexiones, sabiduría y humor
    "¿Cuál dirías que ha sido el momento más valiente de tu vida?",
    "¿Qué época de tu vida te gustaría volver a repetir por unas horas?",
    "¿Cuál es el mejor consejo que te han dado jamás?",
    "¿Qué es lo que más te hace reír hoy en día?",
    "¿Qué manía o costumbre personal tienes desde que eres joven?",
    "¿Qué te da más paz en los momentos de tranquilidad?",
    "¿Qué refrán o dicho popular utilizas más a menudo?",
    "¿Qué virtud valoras más en una persona?",
    "Si pudieras dejar grabado un mensaje para las futuras generaciones de tu familia, ¿qué diría?",
    "¿Qué te hace sentir más orgulloso al mirar atrás en tu vida?",
    "¿Cuál ha sido el mayor susto que te has llevado y en qué quedó?",
    "¿Tienes algún pequeño ritual diario que no perdones nunca?",
    "¿Qué es lo que más te gusta hacer en una tarde tranquila de invierno?",
    "¿Cómo te imaginas que hablarán de ti tus nietos cuando pasen los años?",
    "¿Qué palabra define mejor cómo ha sido tu viaje por la vida hasta ahora?"
]

lista_preguntas = list(BANCO_PREGUNTAS_DEFAULT)

def aplicar_estilo_celda(cell, texto_p1, texto_p2):
    cell.width = Inches(2.5)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="120" w:type="dxa"/>'
        f'<w:left w:w="140" w:type="dxa"/>'
        f'<w:bottom w:w="120" w:type="dxa"/>'
        f'<w:right w:w="140" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)
    
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAFAFA"/>')
    tcPr.append(shd)
    
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="16" w:space="0" w:color="1A365D"/>'
        f'<w:left w:val="single" w:sz="16" w:space="0" w:color="1A365D"/>'
        f'<w:bottom w:val="single" w:sz="16" w:space="0" w:color="1A365D"/>'
        f'<w:right w:val="single" w:sz="16" w:space="0" w:color="1A365D"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    
    run_titulo = p.add_run("✨ RESPONDE ✨\n")
    run_titulo.font.size = Pt(9.5)
    run_titulo.font.bold = True
    run_titulo.font.color.rgb = RGBColor(26, 54, 93)
    
    run_q1 = p.add_run(f"Opción A:\n{texto_p1}\n")
    run_q1.font.size = Pt(9)
    run_q1.font.color.rgb = RGBColor(30, 30, 30)
    
    run_div = p.add_run("— o —\n")
    run_div.font.size = Pt(8.5)
    run_div.font.bold = True
    run_div.font.color.rgb = RGBColor(120, 120, 120)
    
    run_q2 = p.add_run(f"Opción B:\n{texto_p2}")
    run_q2.font.size = Pt(9)
    run_q2.font.color.rgb = RGBColor(30, 30, 30)

def generar_documento(ruta_salida, lista_p):
    doc = Document()
    
    for section in doc.sections:
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.35)
        section.right_margin = Inches(0.35)
        
    preguntas_disponibles = list(lista_p)
    random.shuffle(preguntas_disponibles)
    
    num_preguntas = len(preguntas_disponibles)
    num_tarjetas = num_preguntas // 2
    
    TARJETAS_POR_PAGINA = 12
    idx = 0
    tarjetas_restantes = num_tarjetas
    
    while tarjetas_restantes > 0:
        tarjetas_en_esta_pag = min(TARJETAS_POR_PAGINA, tarjetas_restantes)
        filas_necesarias = (tarjetas_en_esta_pag + 2) // 3
        
        tabla = doc.add_table(rows=filas_necesarias, cols=3)
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        for f in range(filas_necesarias):
            trPr = tabla.rows[f]._tr.get_or_add_trPr()
            trHeight = parse_xml(f'<w:trHeight {nsdecls("w")} w:val="1600" w:hRule="atLeast"/>')
            trPr.append(trHeight)
            
            for c in range(3):
                if idx < num_tarjetas:
                    p1 = preguntas_disponibles[idx * 2]
                    p2 = preguntas_disponibles[idx * 2 + 1]
                    cell = tabla.cell(f, c)
                    aplicar_estilo_celda(cell, p1, p2)
                    idx += 1
                    
        tarjetas_restantes -= tarjetas_en_esta_pag
        if tarjetas_restantes > 0:
            doc.add_page_break()

    doc.save(ruta_salida)

def abrir_carpeta(ruta_archivo):
    carpeta = os.path.dirname(ruta_archivo)
    if platform.system() == "Windows":
        os.startfile(carpeta)
    elif platform.system() == "Darwin":
        subprocess.call(["open", carpeta])
    else:
        subprocess.call(["xdg-open", carpeta])

def iniciar_interfaz():
    def actualizar_contador():
        lbl_count.config(text=f"Total: {len(lista_preguntas)} preguntas ({len(lista_preguntas)//2} tarjetas)")

    def refrescar_listbox():
        lst_preguntas.delete(0, tk.END)
        for i, p in enumerate(lista_preguntas, 1):
            lst_preguntas.insert(tk.END, f"{i}. {p}")
        actualizar_contador()

    def agregar_pregunta():
        p = txt_entrada.get("1.0", tk.END).strip()
        if p:
            lista_preguntas.append(p)
            txt_entrada.delete("1.0", tk.END)
            refrescar_listbox()
        else:
            messagebox.showwarning("Atención", "Escribe una pregunta antes de añadir.")

    def eliminar_pregunta():
        sel = lst_preguntas.curselection()
        if sel:
            idx = sel[0]
            del lista_preguntas[idx]
            refrescar_listbox()
        else:
            messagebox.showwarning("Atención", "Selecciona una pregunta de la lista para eliminar.")

    def modificar_pregunta():
        sel = lst_preguntas.curselection()
        if sel:
            idx = sel[0]
            texto_actual = lista_preguntas[idx]
            
            def guardar_mod():
                nuevo_t = txt_mod.get("1.0", tk.END).strip()
                if nuevo_t:
                    lista_preguntas[idx] = nuevo_t
                    refrescar_listbox()
                    win_mod.destroy()

            win_mod = tk.Toplevel(ventana)
            win_mod.title("Modificar Pregunta")
            win_mod.geometry("400x180")
            win_mod.grab_set()
            
            tk.Label(win_mod, text="Editar pregunta:", font=("Arial", 9, "bold")).pack(anchor="w", padx=10, pady=(10, 2))
            txt_mod = tk.Text(win_mod, height=4, width=45)
            txt_mod.pack(padx=10, pady=5)
            txt_mod.insert("1.0", texto_actual)
            
            tk.Button(win_mod, text="Guardar Cambios", bg="#1A365D", fg="white", font=("Arial", 9, "bold"), command=guardar_mod).pack(pady=5)
        else:
            messagebox.showwarning("Atención", "Selecciona una pregunta para modificar.")

    def cargar_desde_word():
        ruta_doc = filedialog.askopenfilename(
            filetypes=[("Documentos de Word", "*.docx")],
            title="Seleccionar archivo de Word con preguntas"
        )
        if ruta_doc:
            try:
                doc = Document(ruta_doc)
                nuevas_preguntas = []
                
                # Extraer texto de párrafos
                for p in doc.paragraphs:
                    linea = p.text.strip()
                    if linea and len(linea) > 5:
                        nuevas_preguntas.append(linea)
                        
                # Extraer texto de celdas si el documento fuente contenía tablas
                for tabla in doc.tables:
                    for fila in tabla.rows:
                        for celda in fila.cells:
                            linea = celda.text.strip()
                            if linea and len(linea) > 5 and linea not in nuevas_preguntas:
                                nuevas_preguntas.append(linea)
                                
                if nuevas_preguntas:
                    opcion = messagebox.askyesnocancel(
                        "Importar Preguntas",
                        f"Se han encontrado {len(nuevas_preguntas)} preguntas en el documento Word.\n\n"
                        "¿Deseas REEMPLAZAR la lista actual? (Elige 'No' si prefieres AÑADIRLAS a las que ya están)."
                    )
                    
                    if opcion is True: # Reemplazar
                        lista_preguntas.clear()
                        lista_preguntas.extend(nuevas_preguntas)
                        refrescar_listbox()
                        messagebox.showinfo("Éxito", f"Se cargaron {len(nuevas_preguntas)} preguntas sustituyendo a las anteriores.")
                    elif opcion is False: # Añadir / Combinar
                        lista_preguntas.extend(nuevas_preguntas)
                        refrescar_listbox()
                        messagebox.showinfo("Éxito", f"Se añadieron {len(nuevas_preguntas)} preguntas a la lista.")
                else:
                    messagebox.showwarning("Atención", "No se encontraron textos útiles de preguntas en el archivo Word.")
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo leer el archivo Word:\n{str(e)}")

    def restaurar_por_defecto():
        if messagebox.askyesno("Confirmar", "¿Deseas restaurar la lista original de 100 preguntas por defecto?"):
            lista_preguntas.clear()
            lista_preguntas.extend(BANCO_PREGUNTAS_DEFAULT)
            refrescar_listbox()

    def accion_generar():
        if len(lista_preguntas) < 2:
            messagebox.showerror("Error", "Necesitas al menos 2 preguntas para formar una tarjeta.")
            return

        ruta_destino = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Documento de Word", "*.docx")],
            initialfile="Tarjetas_Responde_Importadas.docx",
            title="Guardar tarjetas impresas"
        )
        
        if ruta_destino:
            try:
                generar_documento(ruta_destino, lista_preguntas)
                if messagebox.askyesno("¡Completado!", f"¡Se generaron {len(lista_preguntas)//2} tarjetas con éxito en Word!\n\n¿Deseas abrir la carpeta contenedora ahora?"):
                    abrir_carpeta(ruta_destino)
            except Exception as e:
                messagebox.showerror("Error", f"Ocurrió un error al guardar:\n{str(e)}")

    ventana = tk.Tk()
    ventana.title("Editor y Generador de Tarjetas - Responde")
    ventana.geometry("620x680")
    ventana.resizable(False, False)
    
    ventana.update_idletasks()
    ancho_pantalla = ventana.winfo_screenwidth()
    x = (ancho_pantalla - 620) // 2
    ventana.geometry(f"620x680+{x}+5")

    # Header
    frame_top = tk.Frame(ventana, bg="#1A365D", pady=10)
    frame_top.pack(fill=tk.X)
    
    lbl_title = tk.Label(frame_top, text="Generador de Tarjetas 'RESPONDE'", font=("Arial", 14, "bold"), fg="white", bg="#1A365D")
    lbl_title.pack()

    # Bloque de importación desde Word externo
    frame_import = tk.LabelFrame(ventana, text=" Cargar desde archivo Word externo ", font=("Arial", 9, "bold"), padx=10, pady=5)
    frame_import.pack(fill=tk.X, padx=15, pady=8)

    btn_import = tk.Button(frame_import, text="📂 Seleccionar documento .DOCX con preguntas", font=("Arial", 9, "bold"), bg="#D97706", fg="white", command=cargar_desde_word)
    btn_import.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

    btn_reset = tk.Button(frame_import, text="🔄 Restaurar 100 orig.", font=("Arial", 8), command=restaurar_por_defecto)
    btn_reset.pack(side=tk.RIGHT)
    
    # Bloque para añadir manualmente
    frame_add = tk.LabelFrame(ventana, text=" Añadir nueva pregunta manualmente ", font=("Arial", 9, "bold"), padx=10, pady=5)
    frame_add.pack(fill=tk.X, padx=15, pady=5)
    
    txt_entrada = tk.Text(frame_add, height=2, font=("Arial", 9))
    txt_entrada.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 10))
    
    btn_add = tk.Button(frame_add, text="➕ Añadir", bg="#2E7D32", fg="white", font=("Arial", 9, "bold"), command=agregar_pregunta)
    btn_add.pack(side=tk.RIGHT)

    # Bloque de la lista
    frame_list = tk.LabelFrame(ventana, text=" Banco de Preguntas Activo ", font=("Arial", 9, "bold"), padx=10, pady=5)
    frame_list.pack(fill=tk.BOTH, expand=True, padx=15, pady=5)

    scroll = tk.Scrollbar(frame_list)
    scroll.pack(side=tk.RIGHT, fill=tk.Y)

    lst_preguntas = tk.Listbox(frame_list, yscrollcommand=scroll.set, font=("Arial", 9), selectmode=tk.SINGLE)
    lst_preguntas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    scroll.config(command=lst_preguntas.yview)

    # Botones de edición
    frame_edit_btns = tk.Frame(ventana)
    frame_edit_btns.pack(fill=tk.X, padx=15, pady=5)

    btn_mod = tk.Button(frame_edit_btns, text="✏️️ Modificar seleccionada", font=("Arial", 9), command=modificar_pregunta)
    btn_mod.pack(side=tk.LEFT, padx=(0, 5))

    btn_del = tk.Button(frame_edit_btns, text="🗑️ Eliminar seleccionada", font=("Arial", 9), fg="#C62828", command=eliminar_pregunta)
    btn_del.pack(side=tk.LEFT)

    lbl_count = tk.Label(frame_edit_btns, text="", font=("Arial", 9, "bold"), fg="#1A365D")
    lbl_count.pack(side=tk.RIGHT)

    # Botón principal de guardado
    btn_generar = tk.Button(ventana, text="📄 GENERAR TARJETAS EN WORD (.DOCX)", font=("Arial", 11, "bold"), bg="#1A365D", fg="white", pady=10, command=accion_generar)
    btn_generar.pack(fill=tk.X, padx=15, pady=12)

    refrescar_listbox()
    ventana.mainloop()

if __name__ == "__main__":
    iniciar_interfaz()