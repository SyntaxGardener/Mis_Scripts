import sys
import re
import os
import math
import subprocess
import tkinter as tk
from tkinter import filedialog, messagebox, ttk, colorchooser, scrolledtext
from xml.sax.saxutils import escape

import docx
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, PageBreak, Table, TableStyle
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont


PALETA_DEFECTO = [
    '#2E7D32', '#1565C0', '#C2185B', '#E65100', '#6A1B9A',
    '#00838F', '#D84315', '#455A64', '#8E24AA', '#00695C',
]

FUENTES_ESTANDAR_PDF = [
    'Helvetica', 'Helvetica-Bold', 'Helvetica-Oblique', 'Helvetica-BoldOblique',
    'Times-Roman', 'Times-Bold', 'Times-Italic', 'Times-BoldItalic',
    'Courier', 'Courier-Bold', 'Courier-Oblique', 'Courier-BoldOblique',
    'Symbol', 'ZapfDingbats',
]

SEPARADOR = "==========  NUEVA TARJETA  =========="
MARCA_VACIA = "[TARJETA VACÍA]"
MAX_FILAS_COLS = 10

# Márgenes de seguridad para que la cuadrícula NUNCA rebose la página
MARGEN_SEGURIDAD_X = 4.0
MARGEN_SEGURIDAD_Y = 12.0


# ---------- Utilidades ----------
def es_color_oscuro(hex_color):
    h = hex_color.lstrip('#')
    if len(h) != 6:
        return False
    try:
        r, g, b = int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16)
    except ValueError:
        return False
    return (0.299 * r + 0.587 * g + 0.114 * b) < 128


def limpiar_texto(texto):
    if not texto:
        return ""
    texto = texto.replace('\ufeff', '').replace('\u200b', '')
    texto = re.sub(r'^\s*Tarjeta\s*\d+\s*:\s*', '', texto, flags=re.IGNORECASE)
    texto = texto.strip(' <>«»"')
    texto = re.sub(r'^[<«"]+', '', texto)
    texto = re.sub(r'[>»"]+$', '', texto)
    return texto.strip()


def es_solo_numero(texto):
    if not texto:
        return False
    t = texto.strip()
    if not t or not re.search(r'\d', t):
        return False
    limpio = re.sub(r'[\d\.\s,;:\-–—\)\(\[\]{}]+', '', t)
    if limpio:
        return False
    return len(re.findall(r'\d+', t)) >= 3


def extraer_parrafos_celda(cell):
    partes = [p.text.strip() for p in cell.paragraphs if p.text.strip()]
    return '\n'.join(partes)


def analizar_tablas_docx(ruta_docx):
    doc = docx.Document(ruta_docx)
    info = []
    for i, table in enumerate(doc.tables):
        n_filas = len(table.rows)
        n_cols = len(table.columns) if table.rows else 0
        preview = "(tabla vacía)"
        for row in table.rows:
            for cell in row.cells:
                txt = cell.text.strip()
                if txt:
                    preview = txt.replace('\n', ' ')[:120]
                    break
            if preview != "(tabla vacía)":
                break
        info.append({
            'indice': i,
            'filas': n_filas,
            'columnas': n_cols,
            'preview': preview
        })
    return info


def docx_a_textos(ruta_docx, tablas_seleccionadas=None, incluir_vacias=False):
    doc = docx.Document(ruta_docx)
    textos = []
    if doc.tables:
        for i, table in enumerate(doc.tables):
            if tablas_seleccionadas is not None and i not in tablas_seleccionadas:
                continue
            for row in table.rows:
                for cell in row.cells:
                    bruto = extraer_parrafos_celda(cell)
                    if es_solo_numero(bruto):
                        continue
                    txt = limpiar_texto(bruto)
                    if txt:
                        textos.append(txt)
                    elif incluir_vacias:
                        textos.append("")
    else:
        for p in doc.paragraphs:
            txt = limpiar_texto(p.text)
            if txt and not es_solo_numero(txt):
                textos.append(txt)
    return textos


def _registrar_ttf(ruta, nombre):
    if not os.path.isfile(ruta):
        return False
    try:
        pdfmetrics.registerFont(TTFont(nombre, ruta))
        return True
    except Exception:
        return False


def cargar_fuentes_windows():
    candidatos = {
        'SegoeUI-SemiBoldItalic': r'C:\Windows\Fonts\seguisbi.ttf',
        'SegoeUI-SemiBold':       r'C:\Windows\Fonts\seguisb.ttf',
        'SegoeUI-Italic':         r'C:\Windows\Fonts\segoeuii.ttf',
        'SegoeUI-BoldItalic':     r'C:\Windows\Fonts\segoeuiz.ttf',
        'SegoeUI-Bold':           r'C:\Windows\Fonts\segoeuib.ttf',
        'SegoeUI':                r'C:\Windows\Fonts\segoeui.ttf',
        'Calibri':                r'C:\Windows\Fonts\calibri.ttf',
        'Calibri-Bold':           r'C:\Windows\Fonts\calibrib.ttf',
        'Calibri-Italic':         r'C:\Windows\Fonts\calibrii.ttf',
    }
    cargadas = {}
    for nombre, ruta in candidatos.items():
        if _registrar_ttf(ruta, nombre):
            cargadas[nombre] = ruta
    return cargadas


def abrir_carpeta(ruta_archivo):
    carpeta = os.path.dirname(os.path.abspath(ruta_archivo))
    if not os.path.isdir(carpeta):
        return
    try:
        if sys.platform.startswith('win'):
            os.startfile(carpeta)
        elif sys.platform == 'darwin':
            subprocess.Popen(['open', carpeta])
        else:
            subprocess.Popen(['xdg-open', carpeta])
    except Exception as e:
        print(f"No se pudo abrir la carpeta: {e}")


def dividir_tarjetas(texto_preview):
    patron = re.compile(r'^\s*.*?NUEVA\s+TARJETA.*$',
                        re.IGNORECASE | re.MULTILINE)
    partes = patron.split(texto_preview)
    return [p.strip() for p in partes]


# ---------- Canvas con marcos redondeados ----------
class CanvasTarjetasIndependientes(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.tarjetas_a_dibujar = []

    def registrar_tarjeta(self, x, y, ancho, alto, color):
        self.tarjetas_a_dibujar.append((x, y, ancho, alto, color))

    def showPage(self):
        self.saveState()
        for (x, y, ancho, alto, color) in self.tarjetas_a_dibujar:
            self.setStrokeColor(color)
            self.setLineWidth(3.0)
            self.roundRect(x, y, ancho, alto, radius=14, stroke=1, fill=0)
        self.restoreState()
        self.tarjetas_a_dibujar = []
        super().showPage()


# ---------- Generador de PDF ----------
def generar_pdf(textos, ruta_pdf,
                nombre_fuente='Helvetica', tamano_fuente=11,
                color_texto='#111111',
                colores_marcos=None,
                columnas=2, filas=4):
    if not colores_marcos:
        colores_marcos = ['#2E7D32']
    colores_rl = [colors.HexColor(c) for c in colores_marcos]

    doc = SimpleDocTemplate(
        ruta_pdf, pagesize=A4,
        rightMargin=25, leftMargin=25, topMargin=25, bottomMargin=25
    )
    styles = getSampleStyleSheet()
    estilo_texto = ParagraphStyle(
        'TarjetaTexto', parent=styles['Normal'],
        fontName=nombre_fuente, fontSize=tamano_fuente,
        leading=int(tamano_fuente * 1.45), alignment=1,
        textColor=colors.HexColor(color_texto)
    )

    COLUMNAS = max(1, int(columnas))
    FILAS = max(1, int(filas))
    TARJETAS_POR_PAGINA = COLUMNAS * FILAS

    ANCHO_PAGINA, ALTO_PAGINA = 595.27, 841.89
    MARGEN = 25
    GAP_X = GAP_Y = 15

    # Área útil con margen de seguridad grande
    ANCHO_DISP = ANCHO_PAGINA - 2 * MARGEN - MARGEN_SEGURIDAD_X
    ALTO_DISP = ALTO_PAGINA - 2 * MARGEN - MARGEN_SEGURIDAD_Y

    # Truncamos a 2 decimales hacia abajo para evitar problemas de coma flotante
    ANCHO_TARJETA = math.floor(
        ((ANCHO_DISP - (COLUMNAS - 1) * GAP_X) / COLUMNAS) * 100) / 100
    ALTO_TARJETA = math.floor(
        ((ALTO_DISP - (FILAS - 1) * GAP_Y) / FILAS) * 100) / 100

    def maker(filename, **kw):
        return CanvasTarjetasIndependientes(filename, **kw)

    story = []
    for i in range(0, len(textos), TARJETAS_POR_PAGINA):
        bloque = textos[i:i + TARJETAS_POR_PAGINA]
        celdas = []
        for txt in bloque:
            if txt and txt.strip():
                esc = escape(txt).replace('\n', '<br/>')
                celdas.append(Paragraph(esc, estilo_texto))
            else:
                celdas.append("")
        while len(celdas) < TARJETAS_POR_PAGINA:
            celdas.append("")

        filas_tabla = [[celdas[r * COLUMNAS + c] for c in range(COLUMNAS)]
                       for r in range(FILAS)]
        col_widths = [ANCHO_TARJETA + GAP_X] * (COLUMNAS - 1) + [ANCHO_TARJETA]
        row_heights = [ALTO_TARJETA + GAP_Y] * (FILAS - 1) + [ALTO_TARJETA]

        t = Table(filas_tabla, colWidths=col_widths, rowHeights=row_heights)
        estilos = [
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 10),
            ('LEFTPADDING', (0, 0), (-1, -1), 15),
            ('RIGHTPADDING', (0, 0), (-1, -1), 15),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 10),
        ]
        if COLUMNAS > 1:
            estilos.append(('RIGHTPADDING', (0, 0), (COLUMNAS - 2, -1), 15 + GAP_X))
        if FILAS > 1:
            estilos.append(('BOTTOMPADDING', (0, 0), (-1, FILAS - 2), 10 + GAP_Y))
        t.setStyle(TableStyle(estilos))
        story.append(t)
        if i + TARJETAS_POR_PAGINA < len(textos):
            story.append(PageBreak())

    def dibujar_marcos(canvas_obj, doc_obj):
        page_num = doc_obj.page - 1
        start_idx = page_num * TARJETAS_POR_PAGINA
        for idx_rel in range(TARJETAS_POR_PAGINA):
            idx_global = start_idx + idx_rel
            if idx_global >= len(textos):
                break
            row, col = idx_rel // COLUMNAS, idx_rel % COLUMNAS
            x = MARGEN + col * (ANCHO_TARJETA + GAP_X)
            y = ALTO_PAGINA - MARGEN - (row + 1) * ALTO_TARJETA - row * GAP_Y
            color = colores_rl[idx_global % len(colores_rl)]
            canvas_obj.registrar_tarjeta(x, y, ANCHO_TARJETA, ALTO_TARJETA, color)

    doc.build(story, canvasmaker=maker,
              onFirstPage=dibujar_marcos, onLaterPages=dibujar_marcos)


# ---------- Diálogo de selección de tablas ----------
class DialogoSeleccionTablas(tk.Toplevel):
    def __init__(self, parent, tablas_info):
        super().__init__(parent)
        self.title("Seleccionar tablas del documento")
        self.geometry("760x440")
        self.resultado = None
        self.transient(parent)
        self.grab_set()

        cont = ttk.Frame(self, padding=12)
        cont.pack(fill='both', expand=True)

        ttk.Label(
            cont,
            text=("Tu documento tiene varias tablas. Marca SOLO las que\n"
                  "quieres convertir en tarjetas (por ejemplo, la de los versos,\n"
                  "no la de las soluciones ni la de la lista de números)."),
            justify='left'
        ).pack(anchor='w', pady=(0, 10))

        self.vars = []
        marco_lista = ttk.Frame(cont)
        marco_lista.pack(fill='both', expand=True)

        for info in tablas_info:
            var = tk.BooleanVar(value=True)
            self.vars.append((info['indice'], var))
            texto = (f"Tabla {info['indice'] + 1}   ({info['filas']} filas × "
                     f"{info['columnas']} columnas)\n"
                     f"     {info['preview']}")
            cb = ttk.Checkbutton(marco_lista, text=texto, variable=var)
            cb.pack(anchor='w', pady=3)

        botones = ttk.Frame(cont)
        botones.pack(fill='x', pady=(15, 0))
        ttk.Button(botones, text="Todas", command=self._todas).pack(side='left')
        ttk.Button(botones, text="Ninguna", command=self._ninguna).pack(
            side='left', padx=5)
        ttk.Button(botones, text="Cancelar", command=self._cancelar).pack(
            side='right', padx=5)
        ttk.Button(botones, text="Aceptar", command=self._aceptar).pack(side='right')

    def _todas(self):
        for _, var in self.vars:
            var.set(True)

    def _ninguna(self):
        for _, var in self.vars:
            var.set(False)

    def _aceptar(self):
        self.resultado = {idx for idx, var in self.vars if var.get()}
        self.destroy()

    def _cancelar(self):
        self.resultado = None
        self.destroy()


# ---------- Ventana de previsualización ----------
class VentanaPrevisualizacion(tk.Toplevel):
    def __init__(self, parent, textos, callback_generar):
        super().__init__(parent)
        self.title("Previsualización de tarjetas")
        self.geometry("780x720")
        self.callback_generar = callback_generar
        self.transient(parent)
        self.grab_set()

        cont = ttk.Frame(self, padding=10)
        cont.pack(fill='both', expand=True)

        ttk.Label(
            cont,
            text=("Revisa y edita las tarjetas antes de generar el PDF.\n"
                  f"Cada tarjeta empieza después de una línea '{SEPARADOR}'.\n"
                  f"Las tarjetas sin contenido aparecen como {MARCA_VACIA}.\n"
                  "• Borra una sección para eliminar esa tarjeta.\n"
                  "• Pega el separador para añadir una tarjeta nueva.\n"
                  "• Cambia el orden cortando y pegando bloques."),
            justify='left'
        ).pack(anchor='w', pady=(0, 8))

        self.text = scrolledtext.ScrolledText(cont, wrap='word',
                                              font=('Consolas', 10))
        self.text.pack(fill='both', expand=True)

        self._cargar_textos(textos)

        botones = ttk.Frame(cont)
        botones.pack(fill='x', pady=(10, 0))

        self.lbl_info = ttk.Label(botones, text="", foreground='#666666')
        self.lbl_info.pack(side='left')
        self._actualizar_info()

        ttk.Button(botones, text="Quitar vacías",
                   command=self._quitar_vacias).pack(side='left', padx=(10, 0))
        ttk.Button(botones, text="Cancelar",
                   command=self.destroy).pack(side='right', padx=5)
        ttk.Button(botones, text="Generar PDF",
                   command=self._generar).pack(side='right')

    def _cargar_textos(self, textos):
        formateados = [t if t.strip() else MARCA_VACIA for t in textos]
        contenido = f"\n\n{SEPARADOR}\n\n".join(formateados)
        self.text.delete('1.0', 'end')
        self.text.insert('1.0', contenido)

    def _actualizar_info(self):
        crudo = self.text.get('1.0', 'end-1c')
        partes = dividir_tarjetas(crudo)
        total = len(partes)
        vacias = sum(1 for p in partes if p.strip() == MARCA_VACIA or not p.strip())
        self.lbl_info.config(text=f"{total} tarjetas ({vacias} vacías)")

    def _quitar_vacias(self):
        crudo = self.text.get('1.0', 'end-1c')
        partes = dividir_tarjetas(crudo)
        partes = [p for p in partes if p.strip() and p.strip() != MARCA_VACIA]
        self._cargar_textos(partes)
        self._actualizar_info()

    def _generar(self):
        crudo = self.text.get('1.0', 'end-1c')
        partes = dividir_tarjetas(crudo)
        textos = []
        for p in partes:
            p = p.strip()
            if p == MARCA_VACIA:
                textos.append("")
            else:
                textos.append(p)
        if not any(t.strip() for t in textos):
            messagebox.showwarning("Aviso", "No hay tarjetas para generar.")
            return
        self.callback_generar(textos)
        self.destroy()


# ---------- Aplicación principal ----------
class TarjetasApp:
    def __init__(self, root):
        self.root = root
        root.title("Generador de Tarjetas desde Word")
        root.geometry("760x980")
        root.resizable(False, False)

        self.ruta_docx = tk.StringVar()
        self.ruta_pdf = tk.StringVar()
        self.tamano_fuente = tk.IntVar(value=11)
        self.color_texto = tk.StringVar(value='#111111')
        self.modo_color_marco = tk.StringVar(value='rotativo')
        self.color_unico = tk.StringVar(value='#2E7D32')
        self.columnas = tk.IntVar(value=2)
        self.filas = tk.IntVar(value=7)
        self.info_dist = tk.StringVar()
        self.abrir_al_terminar = tk.BooleanVar(value=True)
        self.incluir_vacias = tk.BooleanVar(value=False)

        self.tablas_seleccionadas = None
        self.tablas_info = []

        self.fuentes_sistema = cargar_fuentes_windows()
        self.fuentes_disponibles = list(FUENTES_ESTANDAR_PDF) + \
                                   list(self.fuentes_sistema.keys())

        if 'SegoeUI-SemiBoldItalic' in self.fuentes_sistema:
            self.fuente_pdf = tk.StringVar(value='SegoeUI-SemiBoldItalic')
        else:
            self.fuente_pdf = tk.StringVar(value='Helvetica')

        self.colores_marcos = list(PALETA_DEFECTO)

        self._construir_ui()
        self._actualizar_distribucion()
        self._actualizar_botones_color()
        self._actualizar_paleta_listbox()

    def _construir_ui(self):
        cont = ttk.Frame(self.root, padding=12)
        cont.pack(fill='both', expand=True)

        # 1) Archivo
        lf1 = ttk.LabelFrame(cont, text=" 1. Archivo Word ", padding=8)
        lf1.pack(fill='x', pady=(0, 8))
        lf1.columnconfigure(0, weight=1)
        ttk.Entry(lf1, textvariable=self.ruta_docx).grid(
            row=0, column=0, sticky='we', padx=(0, 5))
        ttk.Button(lf1, text="Examinar...",
                   command=self._elegir_docx).grid(row=0, column=1)
        ttk.Button(lf1, text="Elegir tablas...",
                   command=self._elegir_tablas).grid(row=0, column=2, padx=(5, 0))
        ttk.Checkbutton(
            lf1,
            text="Mantener tarjetas vacías (celdas que solo tenían 'Tarjeta N:')",
            variable=self.incluir_vacias
        ).grid(row=1, column=0, columnspan=3, sticky='w', pady=(6, 0))

        # 2) Tipografía
        lf2 = ttk.LabelFrame(cont, text=" 2. Tipografía ", padding=8)
        lf2.pack(fill='x', pady=(0, 8))
        lf2.columnconfigure(1, weight=1)
        ttk.Label(lf2, text="Fuente:").grid(row=0, column=0, sticky='w')
        self.combo_fuentes = ttk.Combobox(
            lf2, textvariable=self.fuente_pdf,
            values=self.fuentes_disponibles, state='readonly')
        self.combo_fuentes.grid(row=0, column=1, sticky='we', padx=5)
        ttk.Button(lf2, text="Cargar TTF...",
                   command=self._cargar_ttf).grid(row=0, column=2)
        ttk.Label(lf2, text="Tamaño:").grid(row=1, column=0, sticky='w', pady=(6, 0))
        ttk.Spinbox(lf2, from_=6, to=48, textvariable=self.tamano_fuente,
                    width=6).grid(row=1, column=1, sticky='w', padx=5, pady=(6, 0))
        ttk.Label(lf2, text="Color del texto:").grid(
            row=2, column=0, sticky='w', pady=(6, 0))
        self.btn_color_texto = tk.Button(
            lf2, textvariable=self.color_texto, width=14, relief='groove',
            command=self._elegir_color_texto)
        self.btn_color_texto.grid(row=2, column=1, sticky='w', padx=5, pady=(6, 0))

        # 3) Distribución
        lf3 = ttk.LabelFrame(cont, text=" 3. Distribución por página ", padding=8)
        lf3.pack(fill='x', pady=(0, 8))
        ttk.Label(lf3, text="Columnas:").grid(row=0, column=0, sticky='w')
        ttk.Spinbox(lf3, from_=1, to=MAX_FILAS_COLS, textvariable=self.columnas,
                    width=5, command=self._actualizar_distribucion
                    ).grid(row=0, column=1, sticky='w', padx=(5, 20))
        ttk.Label(lf3, text="Filas:").grid(row=0, column=2, sticky='w')
        ttk.Spinbox(lf3, from_=1, to=MAX_FILAS_COLS, textvariable=self.filas,
                    width=5, command=self._actualizar_distribucion
                    ).grid(row=0, column=3, sticky='w', padx=(5, 20))
        ttk.Label(lf3, textvariable=self.info_dist,
                  foreground='#555555').grid(row=1, column=0, columnspan=6,
                                             sticky='w', pady=(6, 0))

        # 4) Colores de los marcos
        lf4 = ttk.LabelFrame(cont, text=" 4. Colores de los marcos ", padding=8)
        lf4.pack(fill='x', pady=(0, 8))
        lf4.columnconfigure(0, weight=1)
        ttk.Radiobutton(lf4, text="Paleta rotativa",
                        variable=self.modo_color_marco, value='rotativo',
                        command=self._actualizar_botones_color
                        ).grid(row=0, column=0, sticky='w')
        ttk.Radiobutton(lf4, text="Color único:",
                        variable=self.modo_color_marco, value='unico',
                        command=self._actualizar_botones_color
                        ).grid(row=1, column=0, sticky='w', pady=(4, 0))
        self.btn_color_unico = tk.Button(
            lf4, textvariable=self.color_unico, width=14, relief='groove',
            command=self._elegir_color_unico)
        self.btn_color_unico.grid(row=1, column=1, sticky='w', padx=5, pady=(4, 0))

        marco_paleta = ttk.Frame(lf4)
        marco_paleta.grid(row=2, column=0, columnspan=2, sticky='we', pady=(8, 0))
        marco_paleta.columnconfigure(0, weight=1)
        self.lb_paleta = tk.Listbox(marco_paleta, height=5, activestyle='none',
                                    selectmode='extended', exportselection=False)
        self.lb_paleta.grid(row=0, column=0, sticky='we')
        sb = ttk.Scrollbar(marco_paleta, orient='vertical',
                           command=self.lb_paleta.yview)
        sb.grid(row=0, column=1, sticky='ns')
        self.lb_paleta.configure(yscrollcommand=sb.set)
        botones_pal = ttk.Frame(marco_paleta)
        botones_pal.grid(row=0, column=2, sticky='ns', padx=(8, 0))
        ttk.Button(botones_pal, text="Añadir",
                   command=self._añadir_color_paleta).pack(fill='x', pady=(0, 4))
        ttk.Button(botones_pal, text="Quitar",
                   command=self._quitar_color_paleta).pack(fill='x', pady=(0, 4))
        ttk.Button(botones_pal, text="Reiniciar",
                   command=self._reiniciar_paleta).pack(fill='x')

        # 5) Salida
        lf5 = ttk.LabelFrame(cont, text=" 5. Guardar como ", padding=8)
        lf5.pack(fill='x', pady=(0, 8))
        lf5.columnconfigure(0, weight=1)
        ttk.Entry(lf5, textvariable=self.ruta_pdf).grid(
            row=0, column=0, sticky='we', padx=(0, 5))
        ttk.Button(lf5, text="Guardar como...",
                   command=self._elegir_pdf).grid(row=0, column=1)

        ttk.Checkbutton(
            cont, text="Abrir la carpeta al terminar",
            variable=self.abrir_al_terminar
        ).pack(anchor='w', pady=(0, 6))

        ttk.Button(cont, text="Extraer y previsualizar tarjetas",
                   command=self._extraer_y_previsualizar).pack(fill='x', ipady=8)

    # ---------- callbacks ----------
    def _elegir_docx(self):
        ruta = filedialog.askopenfilename(
            title="Seleccionar Word",
            filetypes=[("Archivos Word", "*.docx")])
        if not ruta:
            return
        self.ruta_docx.set(ruta)
        try:
            self.tablas_info = analizar_tablas_docx(ruta)
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo leer el Word:\n{e}")
            self.tablas_info = []
            return
        self.tablas_seleccionadas = None
        if len(self.tablas_info) > 1:
            self._elegir_tablas()
        if not self.ruta_pdf.get():
            base = os.path.splitext(ruta)[0]
            self.ruta_pdf.set(base + "_tarjetas.pdf")

    def _elegir_tablas(self):
        ruta = self.ruta_docx.get().strip()
        if not ruta or not os.path.isfile(ruta):
            messagebox.showwarning("Atención",
                                   "Primero selecciona un archivo Word.")
            return
        if not self.tablas_info:
            try:
                self.tablas_info = analizar_tablas_docx(ruta)
            except Exception as e:
                messagebox.showerror("Error", f"No se pudo leer el Word:\n{e}")
                return
        if not self.tablas_info:
            messagebox.showinfo("Sin tablas",
                                "El documento no contiene tablas.")
            return
        dlg = DialogoSeleccionTablas(self.root, self.tablas_info)
        self.root.wait_window(dlg)
        if dlg.resultado is None:
            return
        self.tablas_seleccionadas = dlg.resultado
        messagebox.showinfo(
            "Tablas seleccionadas",
            f"Se usarán {len(self.tablas_seleccionadas)} tabla(s): "
            + ", ".join(str(i + 1) for i in sorted(self.tablas_seleccionadas)))

    def _elegir_pdf(self):
        ruta = filedialog.asksaveasfilename(
            title="Guardar PDF de Tarjetas",
            defaultextension=".pdf",
            filetypes=[("Archivo PDF", "*.pdf")])
        if ruta:
            self.ruta_pdf.set(ruta)

    def _cargar_ttf(self):
        ruta = filedialog.askopenfilename(
            title="Seleccionar fuente TTF",
            filetypes=[("Fuentes TrueType", "*.ttf"), ("Todos", "*.*")])
        if not ruta:
            return
        base = os.path.splitext(os.path.basename(ruta))[0]
        nombre, i = base, 1
        while nombre in self.fuentes_disponibles:
            i += 1
            nombre = f"{base}{i}"
        if not _registrar_ttf(ruta, nombre):
            messagebox.showerror("Error", "No se pudo registrar la fuente.")
            return
        self.fuentes_disponibles.append(nombre)
        self.combo_fuentes['values'] = self.fuentes_disponibles
        self.fuente_pdf.set(nombre)

    def _elegir_color_texto(self):
        color = colorchooser.askcolor(
            title="Color del texto",
            initialcolor=self.color_texto.get())[1]
        if color:
            self.color_texto.set(color.lower())
            self._actualizar_botones_color()

    def _elegir_color_unico(self):
        color = colorchooser.askcolor(
            title="Color del marco",
            initialcolor=self.color_unico.get())[1]
        if color:
            self.color_unico.set(color.lower())
            self._actualizar_botones_color()

    def _actualizar_botones_color(self):
        c = self.color_texto.get()
        fg = 'white' if es_color_oscuro(c) else 'black'
        self.btn_color_texto.configure(bg=c, fg=fg,
                                       activebackground=c, activeforeground=fg)
        c2 = self.color_unico.get()
        fg2 = 'white' if es_color_oscuro(c2) else 'black'
        self.btn_color_unico.configure(bg=c2, fg=fg2,
                                       activebackground=c2, activeforeground=fg2)
        estado = 'normal' if self.modo_color_marco.get() == 'unico' else 'disabled'
        self.btn_color_unico.configure(state=estado)
        estado_lb = 'normal' if self.modo_color_marco.get() == 'rotativo' else 'disabled'
        self.lb_paleta.configure(state=estado_lb)

    def _actualizar_paleta_listbox(self):
        self.lb_paleta.delete(0, 'end')
        for c in self.colores_marcos:
            self.lb_paleta.insert('end', c.upper())
            idx = self.lb_paleta.size() - 1
            fg = 'white' if es_color_oscuro(c) else 'black'
            self.lb_paleta.itemconfig(idx, background=c, foreground=fg,
                                      selectbackground=c, selectforeground=fg)

    def _añadir_color_paleta(self):
        color = colorchooser.askcolor(title="Añadir color a la paleta")[1]
        if color:
            self.colores_marcos.append(color.lower())
            self._actualizar_paleta_listbox()

    def _quitar_color_paleta(self):
        sel = list(self.lb_paleta.curselection())
        for idx in sorted(sel, reverse=True):
            del self.colores_marcos[idx]
        self._actualizar_paleta_listbox()

    def _reiniciar_paleta(self):
        self.colores_marcos = list(PALETA_DEFECTO)
        self._actualizar_paleta_listbox()

    def _actualizar_distribucion(self):
        try:
            c = max(1, min(MAX_FILAS_COLS, int(self.columnas.get())))
            f = max(1, min(MAX_FILAS_COLS, int(self.filas.get())))
        except (tk.TclError, ValueError):
            return
        self.columnas.set(c)
        self.filas.set(f)
        self.info_dist.set(f"→ {c * f} tarjetas por página ({c} × {f})")

    # ---------- flujo principal ----------
    def _extraer_y_previsualizar(self):
        ruta_docx = self.ruta_docx.get().strip()
        if not ruta_docx or not os.path.isfile(ruta_docx):
            messagebox.showwarning("Atención", "Selecciona un archivo Word válido.")
            return

        try:
            textos = docx_a_textos(
                ruta_docx,
                tablas_seleccionadas=self.tablas_seleccionadas,
                incluir_vacias=self.incluir_vacias.get()
            )
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo leer el Word:\n{e}")
            return

        if not textos:
            messagebox.showwarning("Atención",
                                   "No se encontraron textos con los filtros actuales.\n"
                                   "Prueba a cambiar la selección de tablas.")
            return

        VentanaPrevisualizacion(self.root, textos, self._generar_con_textos)

    def _generar_con_textos(self, textos):
        ruta_pdf = self.ruta_pdf.get().strip()
        if not ruta_pdf:
            ruta_pdf = filedialog.asksaveasfilename(
                title="Guardar PDF de Tarjetas",
                defaultextension=".pdf",
                filetypes=[("Archivo PDF", "*.pdf")])
            if not ruta_pdf:
                return
            self.ruta_pdf.set(ruta_pdf)

        if self.modo_color_marco.get() == 'unico':
            colores = [self.color_unico.get()]
        else:
            colores = list(self.colores_marcos)
            if not colores:
                messagebox.showwarning(
                    "Atención",
                    "La paleta está vacía. Añade al menos un color o usa 'Color único'.")
                return

        try:
            generar_pdf(
                textos, ruta_pdf,
                nombre_fuente=self.fuente_pdf.get(),
                tamano_fuente=self.tamano_fuente.get(),
                color_texto=self.color_texto.get(),
                colores_marcos=colores,
                columnas=self.columnas.get(),
                filas=self.filas.get(),
            )
        except Exception as e:
            messagebox.showerror("Error",
                                 f"Ocurrió un error al crear el PDF:\n{e}")
            return

        if self.abrir_al_terminar.get():
            abrir_carpeta(ruta_pdf)

        messagebox.showinfo(
            "¡Éxito!",
            f"Se han generado {len(textos)} tarjetas en:\n{ruta_pdf}")


# ---------- Punto de entrada ----------
def main():
    root = tk.Tk()
    TarjetasApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()