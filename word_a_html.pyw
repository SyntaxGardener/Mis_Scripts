#!/usr/bin/env pythonw
# -*- coding: utf-8 -*-
"""
Conversor de tabla Word -> HTML para el blog de Educastur.

Requiere la librería python-docx:
    pip install python-docx

Uso: haz doble clic sobre este archivo (word_a_html.pyw) en Windows
y no se abrirá ninguna ventana de consola.
"""

import os
import re
import sys
import subprocess
import tkinter as tk
from tkinter import filedialog, messagebox
from html import escape

_PUNCT_ONLY_RE = re.compile(r'^[\s:;,/\-–—().]+$')

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    root = tk.Tk()
    root.withdraw()
    messagebox.showerror(
        "Falta una librería",
        "Este programa necesita la librería 'python-docx'.\n\n"
        "Abre una terminal (CMD) y ejecuta:\n\n"
        "    pip install python-docx\n\n"
        "y vuelve a abrir este programa."
    )
    sys.exit(1)


# ---------------------------------------------------------------------------
# Lógica de conversión
# ---------------------------------------------------------------------------

def run_text_and_formatting(run_el):
    """Devuelve (texto, negrita, cursiva, subrayado) de un elemento <w:r>."""
    texts = []
    for node in run_el.iter():
        if node.tag == qn('w:t'):
            texts.append(node.text or '')
        elif node.tag == qn('w:br'):
            texts.append('\n')
        elif node.tag == qn('w:tab'):
            texts.append('\t')
    text = ''.join(texts)

    bold = italic = underline = False
    rpr = run_el.find(qn('w:rPr'))
    if rpr is not None:
        bold = rpr.find(qn('w:b')) is not None
        italic = rpr.find(qn('w:i')) is not None
        underline = rpr.find(qn('w:u')) is not None
    return text, bold, italic, underline


def wrap_formatting(text, bold, italic, underline):
    html = escape(text).replace('\n', '<br>')
    if underline:
        html = f'<u>{html}</u>'
    if italic:
        html = f'<em>{html}</em>'
    if bold:
        html = f'<strong>{html}</strong>'
    return html


# Estilos de las "píldoras" de enlaces -----------------------------------

VIDEO_LINK_STYLE = (
    'display:inline-block;background:#e8f1fd;color:#1a5296;padding:3px 10px;'
    'margin:2px 4px 2px 0;border-radius:12px;font-size:12px;font-weight:600;'
    'text-decoration:none;vertical-align:middle;line-height:1.6;'
)

NUMBER_LINK_STYLE = (
    'display:inline-flex;align-items:center;justify-content:center;'
    'width:20px;height:20px;background:#ff9f43;color:#ffffff;'
    'border-radius:50%;font-size:11px;font-weight:700;text-decoration:none;'
    'margin:2px 3px 2px 0;vertical-align:middle;'
)

SEPARATOR_STYLE = 'color:#9aa3ad;font-size:11px;'


def render_link(text, url):
    """Renderiza un hipervínculo como píldora de vídeo o burbuja numerada."""
    if not url:
        return escape(text)
    stripped = text.strip()
    if stripped.isdigit():
        return (
            f'<a href="{escape(url)}" target="_blank" rel="noopener" '
            f'style="{NUMBER_LINK_STYLE}">{escape(stripped)}</a>'
        )
    return (
        f'<a href="{escape(url)}" target="_blank" rel="noopener" '
        f'style="{VIDEO_LINK_STYLE}">&#9654; {escape(text.strip())}</a>'
    )


def paragraph_to_html(paragraph, part):
    """Convierte un párrafo -hipervínculos, formato y texto suelto- a HTML."""
    children = list(paragraph._p)
    has_link = any(c.tag == qn('w:hyperlink') for c in children)
    pieces = []
    for child in children:
        tag = child.tag
        if tag == qn('w:hyperlink'):
            r_id = child.get(qn('r:id'))
            url = ''
            if r_id:
                try:
                    url = part.rels[r_id].target_ref
                except KeyError:
                    url = ''
            text = ''.join(
                run_text_and_formatting(run_el)[0] for run_el in child.iter(qn('w:r'))
            )
            if text.strip():
                pieces.append(render_link(text, url))
        elif tag == qn('w:r'):
            text, bold, italic, underline = run_text_and_formatting(child)
            if not text:
                continue
            formatted = wrap_formatting(text, bold, italic, underline)
            if not has_link:
                # Celda/párrafo sin enlaces (p. ej. una cabecera de sección):
                # se deja tal cual para heredar el color del contenedor.
                pieces.append(formatted)
            elif _PUNCT_ONLY_RE.match(text):
                # Puntuación/conectores sueltos (comas, dos puntos, barras…)
                # entre enlaces: se atenúan para que las píldoras destaquen.
                pieces.append(f'<span style="{SEPARATOR_STYLE}">{formatted}</span>')
            else:
                # Texto con contenido propio (p. ej. "Publicidad:") se
                # mantiene legible en tamaño normal.
                pieces.append(f'<span style="font-size:12.5px;color:#3a3a3a;">{formatted}</span>')
    return ''.join(pieces)


def cell_to_html(cell, part):
    paras = [paragraph_to_html(p, part) for p in cell.paragraphs]
    paras = [p for p in paras if p != '']
    if not paras:
        return '&nbsp;'
    return ''.join(f'<div style="margin:3px 0;">{p}</div>' for p in paras)


def is_section_header_row(row):
    """Una fila cuya única celda con texto es la primera se trata como
    cabecera de sección (p. ej. el nombre de una asignatura)."""
    texts = [c.text.strip() for c in row.cells]
    return len(texts) > 1 and texts[0] != '' and all(t == '' for t in texts[1:])


def table_to_html(table, part, first_row_header):
    rows_html = []
    data_row_index = 0
    for i, row in enumerate(table.rows):
        cells = list(row.cells)

        if is_section_header_row(row):
            content = cell_to_html(cells[0], part)
            style = (
                'background:#2c3e50;color:#ffffff;padding:8px 14px;'
                'font-size:14px;font-weight:700;letter-spacing:.4px;'
                'text-transform:uppercase;border:1px solid #223244;'
            )
            rows_html.append(f'<tr><td colspan="{len(cells)}" style="{style}">{content}</td></tr>')
            continue

        tag = 'th' if (first_row_header and i == 0) else 'td'
        zebra = '#ffffff' if data_row_index % 2 == 0 else '#f5f8fb'
        data_row_index += 1
        cells_html = []
        for j, cell in enumerate(cells):
            content = cell_to_html(cell, part)
            if tag == 'th':
                style = (
                    'border:1px solid #d5dbe3;padding:7px 10px;text-align:left;'
                    'vertical-align:top;background:#eef1f5;font-weight:700;font-size:12.5px;'
                )
            else:
                first_col = (
                    'white-space:nowrap;font-weight:700;color:#2c3e50;text-align:center;'
                    if j == 0 else ''
                )
                style = (
                    f'border:1px solid #e2e6ea;padding:6px 10px;text-align:left;'
                    f'vertical-align:top;font-size:12.5px;background:{zebra};{first_col}'
                )
            cells_html.append(f'<{tag} style="{style}">{content}</{tag}>')
        rows_html.append('<tr>' + ''.join(cells_html) + '</tr>')

    table_style = (
        'border-collapse:collapse;font-family:"Segoe UI",Arial,sans-serif;'
    )
    table_html = f'<table style="{table_style}">\n' + '\n'.join(rows_html) + '\n</table>'

    # Envoltorio para esquinas redondeadas, sombra suave y centrado
    # (el ancho se ajusta al contenido en vez de forzar el 100%).
    wrapper_style = (
        'display:inline-block;border-radius:10px;overflow:hidden;'
        'box-shadow:0 2px 6px rgba(0,0,0,0.15);'
    )
    return (
        f'<div style="text-align:center;">'
        f'<div style="{wrapper_style}">{table_html}</div>'
        f'</div>'
    )


def convert_docx_to_html(docx_path, first_row_header=True):
    doc = Document(docx_path)
    part = doc.part
    if not doc.tables:
        raise ValueError('El documento no contiene ninguna tabla.')
    blocks = []
    for idx, table in enumerate(doc.tables):
        if len(doc.tables) > 1:
            blocks.append(f'<!-- Tabla {idx + 1} -->')
        blocks.append(table_to_html(table, part, first_row_header))
    return '\n\n'.join(blocks)


# ---------------------------------------------------------------------------
# Interfaz gráfica
# ---------------------------------------------------------------------------

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Tabla de Word a HTML — Educastur')
        self.resizable(False, False)
        self.docx_path = None

        width, height = 520, 400
        screen_w = self.winfo_screenwidth()
        x = (screen_w - width) // 2
        y = 5
        self.geometry(f'{width}x{height}+{x}+{y}')
        self.configure(padx=20, pady=16)

        tk.Label(
            self,
            text='Convertir tabla de Word a HTML',
            font=('Segoe UI', 13, 'bold')
        ).pack(pady=(0, 4))

        tk.Label(
            self,
            text='Convierte la tabla de un documento Word (con enlaces incluidos)\n'
                 'en código HTML listo para pegar en el editor del blog de Educastur.',
            font=('Segoe UI', 9),
            justify='center'
        ).pack(pady=(0, 14))

        step1 = tk.LabelFrame(self, text=' 1. Documento Word ', padx=10, pady=10)
        step1.pack(fill='x', pady=6)
        tk.Button(step1, text='Seleccionar archivo .docx…', command=self.select_file).pack(anchor='w')
        self.file_label = tk.Label(
            step1, text='(ningún archivo seleccionado)', fg='#555', wraplength=460, justify='left'
        )
        self.file_label.pack(anchor='w', pady=(6, 0))

        self.header_var = tk.BooleanVar(value=True)
        tk.Checkbutton(
            step1,
            text='La primera fila es la cabecera de la tabla',
            variable=self.header_var
        ).pack(anchor='w', pady=(8, 0))

        step2 = tk.LabelFrame(self, text=' 2. Generar HTML ', padx=10, pady=10)
        step2.pack(fill='x', pady=10)
        self.convert_btn = tk.Button(
            step2, text='Convertir y guardar como…', command=self.convert, state='disabled'
        )
        self.convert_btn.pack(anchor='w')

        self.status_label = tk.Label(self, text='', fg='#1a7a1a', wraplength=460, justify='left')
        self.status_label.pack(pady=(8, 0))

    def select_file(self):
        path = filedialog.askopenfilename(
            title='Selecciona el documento Word con la tabla',
            filetypes=[('Documentos Word', '*.docx')]
        )
        if path:
            self.docx_path = path
            self.file_label.config(text=path)
            self.convert_btn.config(state='normal')
            self.status_label.config(text='')

    def convert(self):
        if not self.docx_path:
            return
        default_name = os.path.splitext(os.path.basename(self.docx_path))[0] + '.html'
        save_path = filedialog.asksaveasfilename(
            title='Guardar HTML como…',
            defaultextension='.html',
            initialfile=default_name,
            filetypes=[('Archivo HTML', '*.html')]
        )
        if not save_path:
            return

        try:
            html = convert_docx_to_html(self.docx_path, self.header_var.get())
        except Exception as e:
            messagebox.showerror('Error al convertir', f'No se pudo convertir el documento:\n\n{e}')
            return

        try:
            with open(save_path, 'w', encoding='utf-8') as f:
                f.write(html)
        except Exception as e:
            messagebox.showerror('Error al guardar', f'No se pudo guardar el archivo:\n\n{e}')
            return

        self.status_label.config(text=f'Archivo guardado correctamente:\n{save_path}')
        self.open_folder(os.path.dirname(save_path))
        messagebox.showinfo('Conversión completada', 'El archivo HTML se ha generado y guardado correctamente.')

    def open_folder(self, folder):
        try:
            if sys.platform.startswith('win'):
                os.startfile(folder)
            elif sys.platform == 'darwin':
                subprocess.run(['open', folder])
            else:
                subprocess.run(['xdg-open', folder])
        except Exception:
            pass


if __name__ == '__main__':
    app = App()
    app.mainloop()
