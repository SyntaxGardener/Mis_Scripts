import tkinter as tk
from tkinter import messagebox, scrolledtext, filedialog
import random
import os

# Intentamos importar las librerías extra
try:
    import openpyxl
    from docx import Document
    EXTRAS_DISPONIBLES = True
except ImportError:
    EXTRAS_DISPONIBLES = False

class SelectorUniversal:
    def __init__(self, root):
        self.root = root
        self.root.title("Selector de Alumnos Universal")
        self.root.geometry("700x750")
        self.root.configure(bg="#2c3e50")
        
        self.lista_completa = []
        self.pendientes = []

        # --- PANEL DE CARGA ---
        self.frame_carga = tk.Frame(root, bg="#34495e", pady=15)
        self.frame_carga.pack(fill="x")

        # Botón para cargar archivo
        self.btn_archivo = tk.Button(self.frame_carga, text="📁 CARGAR (TXT, XLSX, DOCX)", 
                                    command=self.cargar_archivo, bg="#95a5a6", width=35, font=("Arial", 10, "bold"))
        self.btn_archivo.pack(pady=5)

        self.txt_nombres = scrolledtext.ScrolledText(self.frame_carga, height=6, width=50)
        self.txt_nombres.pack(pady=5)
        
        self.btn_confirmar = tk.Button(self.frame_carga, text="✅ CONFIRMAR LISTADO", 
                                      command=self.confirmar_lista, bg="#3498db", fg="white", font=("Arial", 10, "bold"))
        self.btn_confirmar.pack(pady=5)

        # --- VISUALIZACIÓN ---
        self.label_alumno = tk.Label(root, text="ESPERANDO LISTA...", font=("Helvetica", 55, "bold"), 
                                     bg="#2c3e50", fg="#f1c40f", wraplength=650)
        self.label_alumno.pack(expand=True)

        self.label_info = tk.Label(root, text="Alumnos: 0 | Pendientes: 0", fg="#bdc3c7", bg="#2c3e50")
        self.label_info.pack()

        self.btn_sorteo = tk.Button(root, text="🎯 ¡A LA PIZARRA!", command=self.seleccionar,
                                   font=("Arial", 22, "bold"), bg="#27ae60", fg="white", height=2, state="disabled")
        self.btn_sorteo.pack(fill="x", padx=20, pady=20)

    def cargar_archivo(self):
        tipos = [("Todos los soportados", "*.txt *.xlsx *.docx"), 
                 ("Texto", "*.txt"), ("Excel", "*.xlsx"), ("Word", "*.docx")]
        ruta = filedialog.askopenfilename(filetypes=tipos)
        
        if not ruta: return

        nombres_extraidos = []
        ext = os.path.splitext(ruta)[1].lower()

        try:
            if ext == ".txt":
                with open(ruta, 'r', encoding='utf-8') as f:
                    nombres_extraidos = [line.strip() for line in f if line.strip()]
            
            elif ext == ".xlsx":
                if not EXTRAS_DISPONIBLES: raise ImportError
                wb = openpyxl.load_workbook(ruta, data_only=True)
                hoja = wb.active
                for row in hoja.iter_rows(min_col=1, max_col=1, values_only=True):
                    if row[0]: nombres_extraidos.append(str(row[0]).strip())
            
            elif ext == ".docx":
                if not EXTRAS_DISPONIBLES: raise ImportError
                doc = Document(ruta)
                nombres_extraidos = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

            # Volcar al cuadro de texto para que el profe lo vea
            self.txt_nombres.delete("1.0", tk.END)
            self.txt_nombres.insert(tk.END, "\n".join(nombres_extraidos))
            messagebox.showinfo("Carga", f"Se han detectado {len(nombres_extraidos)} nombres.")

        except ImportError:
            messagebox.showerror("Librerías faltantes", "Para leer Excel/Word necesitas instalar:\npip install openpyxl python-docx")
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo leer el archivo: {e}")

    def confirmar_lista(self):
        contenido = self.txt_nombres.get("1.0", tk.END).strip()
        self.lista_completa = [n.strip() for n in contenido.split('\n') if n.strip()]
        self.pendientes = list(self.lista_completa)
        random.shuffle(self.pendientes)
        self.btn_sorteo.config(state="normal")
        self.actualizar_stats()

    def actualizar_stats(self):
        self.label_info.config(text=f"Total: {len(self.lista_completa)} | Pendientes: {len(self.pendientes)}")

    def seleccionar(self):
        if not self.pendientes:
            if messagebox.askyesno("Fin", "¿Reiniciar lista?"):
                self.pendientes = list(self.lista_completa)
                random.shuffle(self.pendientes)
            else: return
        self.label_alumno.config(text=self.pendientes.pop(), fg="#1abc9c")
        self.actualizar_stats()

if __name__ == "__main__":
    root = tk.Tk()
    app = SelectorUniversal(root)
    root.mainloop()