#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Formateador de preguntas tipo test
----------------------------------
Convierte preguntas con formato:
    ●   1. ¿Pregunta? a) opción  b) opción  c) opción  d) opción

En un documento Word (.docx) con la pregunta y las opciones (A, B, C, D)
sangradas respecto a ella, organizadas en 1, 2 o 4 por línea según su
longitud.

Guardar como .pyw para que no se abra una consola al ejecutarlo con doble clic.
"""

import os
import re
import sys

# --------------------------------------------------------------------------
# Ocultar la consola si, por el motivo que sea (asociación de archivos,
# lanzador externo, etc.), Python se ejecutó con consola visible.
# --------------------------------------------------------------------------
if sys.platform.startswith('win'):
    try:
        import ctypes
        _hwnd = ctypes.windll.kernel32.GetConsoleWindow()
        if _hwnd:
            ctypes.windll.user32.ShowWindow(_hwnd, 0)  # SW_HIDE
    except Exception:
        pass

import subprocess
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinter import font as tkfont

try:
    from docx import Document
    from docx.shared import Cm, Pt
except ImportError:
    Document = None


# --------------------------------------------------------------------------
# Configuración de fuente
# --------------------------------------------------------------------------

def elegir_fuente(preferida=("Calibri", "Arial")):
    """Devuelve la primera fuente disponible en el sistema de la lista dada."""
    disponibles = set(tkfont.families())
    for nombre in preferida:
        if nombre in disponibles:
            return nombre
    return "TkDefaultFont"


SANGRIA_PREVIEW = "     "  # sangría (en la vista previa de texto) de opciones respecto a la pregunta
SANGRIA_DOCX_CM = 1.0      # sangría (en el .docx) en centímetros

# --------------------------------------------------------------------------
# Lógica de conversión
# --------------------------------------------------------------------------

PATRON = re.compile(
    r'^\s*[●•\-\*]?\s*'
    r'(?P<num>\d+)\.\s*'
    r'(?P<pregunta>.+?)\s+'
    r'a\)\s*(?P<a>.+?)\s+'
    r'b\)\s*(?P<b>.+?)\s+'
    r'c\)\s*(?P<c>.+?)\s+'
    r'd\)\s*(?P<d>.+?)\s*$',
    re.IGNORECASE
)


def analizar_texto(texto_original):
    """
    Recorre el texto línea a línea y devuelve una lista de bloques:
      ('pregunta', num, pregunta, [op_a, op_b, op_c, op_d])
      ('texto', linea_original)                                  <- si no se reconoce
    """
    bloques = []
    for linea in texto_original.splitlines():
        linea_limpia = linea.strip()
        if not linea_limpia:
            continue

        m = PATRON.match(linea_limpia)
        if m:
            num = m.group('num')
            pregunta = m.group('pregunta').strip()
            opciones = [
                m.group('a').strip(),
                m.group('b').strip(),
                m.group('c').strip(),
                m.group('d').strip(),
            ]
            bloques.append(('pregunta', num, pregunta, opciones))
        else:
            bloques.append(('texto', linea_limpia))

    return bloques


def vista_previa_texto(bloques):
    """Genera el texto de vista previa (para el cuadro de resultado en pantalla)."""
    letras = ['A', 'B', 'C', 'D']
    lineas_totales = []

    for bloque in bloques:
        if bloque[0] == 'texto':
            lineas_totales.append(bloque[1])
            continue

        _, num, pregunta, opciones = bloque
        max_len = max(len(o) for o in opciones)
        lineas_totales.append(f"{num}. {pregunta}")

        if max_len <= 10:
            ancho_col = max_len + 20
            partes = [f"{l}.  {o}".ljust(ancho_col) for l, o in zip(letras, opciones)]
            lineas_totales.append(SANGRIA_PREVIEW + "".join(partes).rstrip())
        elif max_len <= 32:
            ancho_col = max_len + 24
            l1 = SANGRIA_PREVIEW + f"A.  {opciones[0]}".ljust(ancho_col) + f"B.  {opciones[1]}"
            l2 = SANGRIA_PREVIEW + f"C.  {opciones[2]}".ljust(ancho_col) + f"D.  {opciones[3]}"
            lineas_totales.append(l1)
            lineas_totales.append(l2)
        else:
            for l, o in zip(letras, opciones):
                lineas_totales.append(SANGRIA_PREVIEW + f"{l}.  {o}")

    return "\n".join(lineas_totales)


def generar_docx(bloques, ruta, fuente_nombre):
    """Crea el documento Word con la pregunta y las opciones sangradas."""
    if Document is None:
        raise RuntimeError(
            "Falta la librería 'python-docx'. Instálala con:\n\n    pip install python-docx"
        )

    doc = Document()

    estilo_normal = doc.styles['Normal']
    estilo_normal.font.name = fuente_nombre
    estilo_normal.font.size = Pt(11)

    letras = ['A', 'B', 'C', 'D']

    for bloque in bloques:
        if bloque[0] == 'texto':
            doc.add_paragraph(bloque[1])
            continue

        _, num, pregunta, opciones = bloque
        max_len = max(len(o) for o in opciones)

        p_pregunta = doc.add_paragraph(f"{num}. {pregunta}")
        p_pregunta.paragraph_format.space_after = Pt(2)
        p_pregunta.paragraph_format.space_before = Pt(6)

        indent = Cm(SANGRIA_DOCX_CM)

        # Ancho de columna calculado a partir del contenido real (letra + opción),
        # con un margen de separación pequeño para que las tabulaciones no queden
        # excesivamente separadas.
        CHAR_CM = 0.16      # ancho aproximado de un carácter en Calibri/Arial 11pt
        PREFIJO_CHARS = 4   # "A.  "
        GAP_CM = 0.55       # separación adicional entre columnas
        ancho_col = Cm((PREFIJO_CHARS + max_len) * CHAR_CM + GAP_CM)

        if max_len <= 10:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = indent
            p.paragraph_format.space_after = Pt(0)
            for i, (letra, op) in enumerate(zip(letras, opciones)):
                if i > 0:
                    p.add_run("\t")
                p.add_run(f"{letra}.  {op}")
            for i in range(1, 4):
                p.paragraph_format.tab_stops.add_tab_stop(Cm(indent.cm + ancho_col.cm * i))

        elif max_len <= 32:
            for l1_idx, l2_idx in [(0, 1), (2, 3)]:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = indent
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(indent.cm + ancho_col.cm))
                p.add_run(f"{letras[l1_idx]}.  {opciones[l1_idx]}\t{letras[l2_idx]}.  {opciones[l2_idx]}")

        else:
            for letra, op in zip(letras, opciones):
                p = doc.add_paragraph(f"{letra}.  {op}")
                p.paragraph_format.left_indent = indent
                p.paragraph_format.space_after = Pt(0)

    doc.save(ruta)


def abrir_carpeta(ruta_archivo):
    """Abre el explorador de archivos en la carpeta que contiene el archivo."""
    carpeta = os.path.dirname(os.path.abspath(ruta_archivo))
    try:
        if sys.platform.startswith('win'):
            os.startfile(carpeta)
        elif sys.platform == 'darwin':
            subprocess.Popen(['open', carpeta])
        else:
            subprocess.Popen(['xdg-open', carpeta])
    except Exception:
        pass


# --------------------------------------------------------------------------
# Interfaz gráfica
# --------------------------------------------------------------------------

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Formateador de preguntas tipo test")

        self.fuente_nombre = elegir_fuente()
        self.fuente_texto = (self.fuente_nombre, 11)
        self.fuente_negrita = (self.fuente_nombre, 10, "bold")

        self._dimensionar_y_centrar()

        # ---- Entrada (a 5 píxeles del borde superior) ----
        self.texto_entrada = tk.Text(self, wrap="word", font=self.fuente_texto)
        self.texto_entrada.pack(fill=tk.BOTH, expand=True, padx=10, pady=(5, 0))

        tk.Label(
            self,
            text="↑ Texto original (pega aquí las preguntas)",
            font=(self.fuente_nombre, 8), fg="gray40"
        ).pack(anchor="w", padx=10)

        # ---- Botones ----
        marco_botones = tk.Frame(self)
        marco_botones.pack(fill=tk.X, padx=10, pady=5)

        ttk.Button(
            marco_botones, text="Pegar", command=self.pegar
        ).pack(side=tk.LEFT)

        ttk.Button(
            marco_botones, text="Convertir formato", command=self.convertir
        ).pack(side=tk.LEFT, padx=(8, 0))

        ttk.Button(
            marco_botones, text="Guardar como .docx...", command=self.guardar
        ).pack(side=tk.LEFT, padx=(8, 0))

        ttk.Button(
            marco_botones, text="Limpiar todo", command=self.limpiar
        ).pack(side=tk.LEFT, padx=(8, 0))

        self.etiqueta_estado = tk.Label(marco_botones, text="", fg="gray30", font=(self.fuente_nombre, 9))
        self.etiqueta_estado.pack(side=tk.LEFT, padx=(15, 0))

        # ---- Salida ----
        marco_salida = tk.Frame(self)
        marco_salida.pack(fill=tk.BOTH, expand=True, padx=10, pady=(0, 10))

        tk.Label(
            marco_salida,
            text="Resultado convertido:",
            font=self.fuente_negrita
        ).pack(anchor="w")

        self.texto_salida = tk.Text(marco_salida, wrap="none", font=self.fuente_texto)
        self.texto_salida.pack(fill=tk.BOTH, expand=True, pady=(3, 0))

        self.bloques_actuales = []  # resultado del último análisis, listo para exportar a docx
        self.texto_entrada.focus_set()

    # ------------------------------------------------------------------
    def _dimensionar_y_centrar(self):
        self.update_idletasks()
        pantalla_ancho = self.winfo_screenwidth()
        pantalla_alto = self.winfo_screenheight()

        # La ventana nunca ocupa más del 85% del alto/ancho de pantalla disponibles
        ancho = min(900, int(pantalla_ancho * 0.85))
        alto = min(700, int(pantalla_alto * 0.85))

        x = (pantalla_ancho - ancho) // 2
        y = 5  # la ventana queda a 5 píxeles del borde superior de la pantalla

        self.geometry(f"{ancho}x{alto}+{x}+{y}")
        self.minsize(min(600, ancho), min(450, alto))

    def pegar(self):
        try:
            contenido = self.clipboard_get()
        except tk.TclError:
            messagebox.showwarning("Aviso", "El portapapeles está vacío o no contiene texto.")
            return
        self.texto_entrada.focus_set()
        self.texto_entrada.insert(tk.INSERT, contenido)

    def convertir(self):
        original = self.texto_entrada.get("1.0", tk.END)
        if not original.strip():
            messagebox.showwarning("Aviso", "No hay texto para convertir.")
            return

        self.bloques_actuales = analizar_texto(original)
        resultado = vista_previa_texto(self.bloques_actuales)

        self.texto_salida.delete("1.0", tk.END)
        self.texto_salida.insert("1.0", resultado)

        sin_reconocer = sum(1 for b in self.bloques_actuales if b[0] == 'texto')
        if sin_reconocer:
            self.etiqueta_estado.config(
                text=f"Convertido. {sin_reconocer} línea(s) no reconocida(s) se dejaron igual.",
                fg="darkorange"
            )
        else:
            self.etiqueta_estado.config(text="Convertido correctamente.", fg="darkgreen")

    def guardar(self):
        if not self.bloques_actuales:
            messagebox.showwarning("Aviso", "No hay resultado que guardar. Convierte primero el texto.")
            return

        ruta = filedialog.asksaveasfilename(
            title="Guardar documento Word",
            defaultextension=".docx",
            initialfile="preguntas_formateadas.docx",
            filetypes=[("Documento de Word", "*.docx"), ("Todos los archivos", "*.*")]
        )
        if not ruta:
            return

        nombre_base = os.path.splitext(os.path.basename(ruta))[0].strip()
        if not nombre_base:
            messagebox.showwarning("Aviso", "El nombre de archivo no puede estar en blanco.")
            return

        try:
            generar_docx(self.bloques_actuales, ruta, self.fuente_nombre)
        except Exception as e:
            messagebox.showerror("Error", f"No se pudo guardar el archivo:\n{e}")
            return

        abrir_carpeta(ruta)
        self.lift()
        self.focus_force()

        # Deja la aplicación lista para pegar y convertir un nuevo texto,
        # sin necesidad de cerrarla y volver a abrirla.
        self.texto_entrada.delete("1.0", tk.END)
        self.texto_salida.delete("1.0", tk.END)
        self.bloques_actuales = []
        self.etiqueta_estado.config(text=f"Guardado en: {ruta}  ·  Listo para un nuevo texto.", fg="darkgreen")
        self.texto_entrada.focus_set()

    def limpiar(self):
        self.texto_entrada.delete("1.0", tk.END)
        self.texto_salida.delete("1.0", tk.END)
        self.bloques_actuales = []
        self.etiqueta_estado.config(text="")
        self.texto_entrada.focus_set()


if __name__ == "__main__":
    app = App()
    app.mainloop()
